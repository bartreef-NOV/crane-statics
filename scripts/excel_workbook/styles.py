"""Cell styles shared by the workbook writers: input vs formula colour coding."""

from __future__ import annotations

from typing import Any

from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

#: Light yellow: a value the user types.
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
#: Light blue: a live formula (do not overwrite).
FORMULA_FILL = PatternFill("solid", fgColor="DDEBF7")
#: Grey: written by the generator, not an input (e.g. overflow counts).
GENERATED_FILL = PatternFill("solid", fgColor="EDEDED")
#: Light red: must stay blank (table guard rows).
GUARD_FILL = PatternFill("solid", fgColor="FCE4E4")
HEADER_FILL = PatternFill("solid", fgColor="D9D9D9")

TITLE_FONT = Font(bold=True, size=14)
SECTION_FONT = Font(bold=True, size=12)
HEADER_FONT = Font(bold=True)
KEY_FONT = Font(italic=True, color="808080", size=9)
NOTE_FONT = Font(italic=True, color="595959")


def put(ws: Worksheet, coord: str, value: Any, fill: PatternFill | None = None) -> None:
    """Write ``value`` at ``coord``; formulas (``=...``) are always formula-coloured."""
    c = ws[coord]
    c.value = value
    if isinstance(value, str) and value.startswith("="):
        c.fill = FORMULA_FILL
    elif fill is not None:
        c.fill = fill


def legend(ws: Worksheet, row: int) -> None:
    """The colour legend on one row (columns A-D)."""
    for col, (text, fill) in zip(
        "ABCD",
        (
            ("Input (type here)", INPUT_FILL),
            ("Formula (do not edit)", FORMULA_FILL),
            ("Generator marker", GENERATED_FILL),
            ("Guard row (keep blank)", GUARD_FILL),
        ),
    ):
        put(ws, f"{col}{row}", text, fill)
