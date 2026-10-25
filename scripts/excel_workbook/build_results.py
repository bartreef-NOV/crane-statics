"""The Results sheet: rope tensions per hoist and the support reactions, as the Scene shows.

Sources:

* rope rows -- ``core/analysis_services.py`` ``compute_result_rows``: one row per hoist rope,
  the solve rope (published as channel ``boom_hoist``) reading the solved tension, the
  others ``baseline + sensitivity x hook load`` (= the Loads rope tension: hook load / parts
  for the scene hoist, hook mass x (|g| + az') / parts for idle hoists).
* reactions -- ``calculation/solver/crane_solver.py`` ``_reactions_from_raw`` (the
  ``_run_statics`` outputs as Fx/Fz/My per support), ``core/demand_extraction.py``
  ``extract_connection_component_response`` (resultant) and ``calculation/engine.py``
  ``resolve_connection_component_linear_response`` (leg_axial / leg_normal).
* out-of-plane rows (heel_left/right, platform base Fy, Mx and the radial / bearing rows) --
  ``build_oop.oop_component_formula`` (OutOfPlane sheet).

When Statics has no solution or the boom width is not > 0, every value cell shows the short
reason (text), never an error value; spare hoist rows read 0 and "".
"""

from __future__ import annotations

from collections.abc import Callable

from openpyxl.worksheet.worksheet import Worksheet

from . import styles
from .build_oop import oop_component_formula
from .layout import RESULT_ROWS, RESULTS_NOTES, Block, Layout, Table, split


def _local(address: str) -> str:
    return split(address)[1]


def _status_formulas(lay: Layout) -> dict[str, str]:
    st = lambda key: lay.statics_cell("statics_status", key)  # noqa: E731
    me = lambda key: lay.results_cell("results_status", key)  # noqa: E731
    width_ok = lay.oop.cell("oop_scene", "width_ok")
    return {
        "ok": f"=IF(AND({st('ok')}=1,{width_ok}=1),1,0)",
        "reason": (
            f'=IF({me("ok")}=1,"",IF({st("solve_ropes")}=0,"no solve rope",'
            f'IF(OR({st("boom_singular")}=1,{st("leg_singular")}=1),"singular matrix",'
            f'IF({width_ok}<>1,"boom width must be > 0","no solution"))))'
        ),
        "message": (
            f'=IF({me("ok")}=1,"OK",IF({st("solve_ropes")}=0,"No solve rope: no active rope '
            f"has role solve (Inputs ropes), so the boom matrix is singular and the app shows "
            f'no results",IF(OR({st("boom_singular")}=1,{st("leg_singular")}=1),"Singular '
            f"matrix on Statics (boom 3x3 or leg frame 6x6): the app finds no solution and "
            f'shows no results",IF({width_ok}<>1,"Boom width (Inputs structure) must be > 0: '
            f"the app rejects it and its out-of-plane split divides by it, so it shows no "
            f'results","No solution"))))'
        ),
    }


def _tension_row(lay: Layout, i: int) -> dict[str, str]:
    inp = lambda key: lay.addr("hoists", key, i)  # noqa: E731
    me = lambda key: lay.results_addr("tensions", key, i)  # noqa: E731
    rt = lambda key: lay.loads_col_range("rope_tensions", key)  # noqa: E731
    on = f"{me('active')}=1"
    found = f'MATCH({me("rope_id")},{rt("id_key")},0)'
    return {
        "active": f"={inp('active')}",
        "hoist_id": f'=IF({on},{inp("id")}&"","")',
        "rope_id": f'=IF({on},{inp("rope_id")}&"","")',
        "role": f'=IF({me("rope_id")}="","",IFERROR(INDEX({rt("role")},{found})&"",""))',
        "tension_kN": (
            f"=IF({me('active')}<>1,0,IF({lay.results_cell('results_status', 'ok')}<>1,"
            f"{lay.results_cell('results_status', 'reason')},"
            f'IF({me("role")}="solve",{lay.statics_value("t_bh")},'
            f'IF({me("role")}="",0,INDEX({rt("tension_kN")},{found})))))'
        ),
    }


def _component_formula(lay: Layout, entity: str, component: str) -> str:
    """The number behind ``reaction.<entity>.<component>`` (no leading ``=``)."""
    out_of_plane = oop_component_formula(lay, entity, component)
    if out_of_plane is not None:
        return out_of_plane
    v = lay.statics_value
    raw = {
        "heel_point": ("h_x", "h_z"),
        "rear_hinge": ("r_rb_x", "r_rb_z"),
        "front_hinge": ("r_fb_x", "r_fb_z"),
        "upper_hinge": ("r_top_x", "r_top_z"),
        "platform_base": ("r_g_x", "r_g_z"),
    }
    fx, fz = (v(key) for key in raw[entity])
    if component == "fx":
        return fx
    if component == "fz":
        return fz
    if component == "my":
        return v("m_g")
    if component == "resultant":
        return f"SQRT({fx}^2+{fz}^2)"
    if component in ("leg_axial", "leg_normal"):
        leg = {"rear_hinge": "rear", "front_hinge": "front"}[entity]
        ux = lay.statics_cell("leg_axes", f"{leg}_ux")
        uz = lay.statics_cell("leg_axes", f"{leg}_uz")
        ok = lay.statics_cell("leg_axes", f"{leg}_axis_ok")
        value = f"{fx}*{ux}+{fz}*{uz}" if component == "leg_axial" else f"-{fx}*{uz}+{fz}*{ux}"
        return f'IF({ok}=1,{value},"n/a (zero-length leg)")'
    raise KeyError(f"no formula for {entity}.{component}")


def _reaction_row(lay: Layout, i: int) -> dict[str, str]:
    entity, component, unit, source = RESULT_ROWS[i - 1]
    status = lambda key: lay.results_cell("results_status", key)  # noqa: E731
    return {
        "key": f"reaction.{entity}.{component}",
        "entity": entity,
        "component": component,
        "value": (
            f"=IF({status('ok')}<>1,{status('reason')},"
            f"{_component_formula(lay, entity, component)})"
        ),
        "unit": unit,
        "source": source,
    }


_ROW_FORMULAS: dict[str, Callable[[Layout, int], dict[str, str]]] = {
    "tensions": _tension_row,
    "reactions": _reaction_row,
}


def _write_block(ws: Worksheet, block: Block, formulas: dict[str, str]) -> None:
    ws[f"A{block.title_row}"] = block.title
    ws[f"A{block.title_row}"].font = styles.SECTION_FONT
    for f in block.fields:
        row = block.row(f.key)
        ws[f"A{row}"] = f.header
        styles.put(ws, f"B{row}", formulas[f.key])
        ws[f"C{row}"] = f.unit or None
        ws[f"D{row}"] = f.key + (f"  -- {f.note}" if f.note else "")
        ws[f"D{row}"].font = styles.KEY_FONT


def _write_table(ws: Worksheet, lay: Layout, table: Table) -> None:
    ws[f"A{table.title_row}"] = table.title
    ws[f"A{table.title_row}"].font = styles.SECTION_FONT
    for f in table.fields:
        col = table.column(f.key)
        header = ws[f"{col}{table.header_row}"]
        header.value = f"{f.header} [{f.unit}]" if f.unit else f.header
        header.font = styles.HEADER_FONT
        header.fill = styles.HEADER_FILL
        key = ws[f"{col}{table.key_row}"]
        key.value = f"{f.key} ({f.note})" if f.note else f.key
        key.font = styles.KEY_FONT
    rows = _ROW_FORMULAS[table.name]
    for i in range(1, table.capacity + 1):
        for key, value in rows(lay, i).items():
            styles.put(ws, _local(lay.results_addr(table.name, key, i)), value)


def write_results(ws: Worksheet, lay: Layout) -> None:
    """Fill the Results sheet (all formulas, driven by Statics and Loads)."""
    ws["A1"] = "Results -- rope tensions and support reactions"
    ws["A1"].font = styles.TITLE_FONT
    for k, note in enumerate(RESULTS_NOTES):
        ws[f"A{3 + k}"] = note
        ws[f"A{3 + k}"].font = styles.NOTE_FONT
    _write_block(ws, lay.results_blocks["results_status"], _status_formulas(lay))
    for table in lay.results_tables.values():
        _write_table(ws, lay, table)
    ws.column_dimensions["A"].width = 44
    for col in "BCDEF":
        ws.column_dimensions[col].width = 22
