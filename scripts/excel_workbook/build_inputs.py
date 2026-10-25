"""The Inputs sheet: single-value blocks and fixed-capacity tables from ``expanded_model``."""

from __future__ import annotations

import dataclasses
from typing import Any

from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from . import styles
from .layout import FIXED_NODE_POINTS, Block, Field, Layout, Table, split
from .reference import SceneCase


def table_rows(model: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    """The model's rows per Inputs table (keys are the table field keys)."""
    return {
        "mass_points": [
            {"body": b["id"], **mp} for b in model["bodies"] for mp in b["mass_points"]
        ],
        "wind_areas": [{"body": b["id"], **wa} for b in model["bodies"] for wa in b["wind_areas"]],
        "nodes": list(model["nodes"]),
        "segments": [{"rope_id": r["id"], **s} for r in model["ropes"] for s in r["segments"]],
        "ropes": [{"id": r["id"], "role": r["tension_mode"]} for r in model["ropes"]],
        "hoists": list(model["hoists"]),
    }


def block_values(model: dict[str, Any], case: SceneCase) -> dict[str, dict[str, Any]]:
    env = model["environment"]
    structure = model["structure"]
    points = structure["points"]
    return {
        "environment": {
            "gravity_m_s2": env["gravity_m_s2"],
            "air_density_kg_m3": env["air_density_kg_m3"],
            "max_wind_speed_m_s": env["max_wind_speed_m_s"],
            "reference_x_m": env["vessel_motion"]["reference_x_m"],
            "reference_z_m": env["vessel_motion"]["reference_z_m"],
        },
        "structure": {
            "boom_width_m": structure["boom_width_m"],
            "slewing_bearing_diameter_m": structure["slewing_bearing_diameter_m"],
            **{
                f"{point}_{axis}": values[axis]
                for point, values in points.items()
                for axis in ("x_m", "z_m")
            },
        },
        "scene": dataclasses.asdict(case),
    }


def _cell_value(value: Any) -> Any:
    if isinstance(value, bool):
        return 1 if value else 0
    return value


def _validation(ws: Worksheet, f: Field, sqref: str) -> None:
    choices = ("1", "0") if f.kind == "flag" else f.choices
    if not choices:
        return
    dv = DataValidation(
        type="list",
        formula1='"' + ",".join(choices) + '"',
        allow_blank=True,
        showErrorMessage=True,
        errorStyle="stop",
        errorTitle="Not an allowed value",
        error="Pick one of: " + ", ".join(choices),
    )
    dv.add(sqref)
    ws.add_data_validation(dv)


def _write_block(ws: Worksheet, block: Block, values: dict[str, Any]) -> None:
    ws[f"A{block.title_row}"] = block.title
    ws[f"A{block.title_row}"].font = styles.SECTION_FONT
    for f in block.fields:
        row = block.row(f.key)
        ws[f"A{row}"] = f.header
        styles.put(ws, f"B{row}", _cell_value(values.get(f.key)), styles.INPUT_FILL)
        ws[f"C{row}"] = f.unit or None
        ws[f"D{row}"] = f"{f.key}" + (f"  -- {f.note}" if f.note else "")
        ws[f"D{row}"].font = styles.KEY_FONT
        _validation(ws, f, f"B{row}")


def _write_table(ws: Worksheet, lay: Layout, table: Table, rows: list[dict[str, Any]]) -> None:
    t = table
    ws[f"A{t.title_row}"] = t.title
    ws[f"A{t.title_row}"].font = styles.SECTION_FONT
    ws[f"D{t.title_row}"] = "Capacity"
    styles.put(ws, split(t.capacity_addr)[1], t.capacity, styles.GENERATED_FILL)
    ws[f"F{t.title_row}"] = "Rows not written (overflow)"
    styles.put(ws, split(t.overflow_addr)[1], max(0, len(rows) - t.capacity), styles.GENERATED_FILL)
    for f in t.fields:
        col = t.column(f.key)
        header = ws[f"{col}{t.header_row}"]
        header.value = f"{f.header} [{f.unit}]" if f.unit else f.header
        header.font = styles.HEADER_FONT
        header.fill = styles.HEADER_FILL
        key = ws[f"{col}{t.key_row}"]
        key.value = f.key
        key.font = styles.KEY_FONT
        if f.note:
            key.value = f"{f.key} ({f.note})"
        _validation(ws, f, f"{col}{t.first_row}:{col}{t.last_row}")
        ws[f"{col}{t.guard_row}"].fill = styles.GUARD_FILL

    for i in range(1, t.capacity + 1):
        row = rows[i - 1] if i <= len(rows) else None
        for f in t.fields:
            coord = split(t.addr(f.key, i))[1]
            if row is None:
                value: Any = 0 if f.key == "active" else None
            elif f.key == "active":
                value = 1
            else:
                value = _cell_value(row.get(f.key))
                if t.name == "nodes" and f.key in ("x_m", "z_m") and row["id"] in FIXED_NODE_POINTS:
                    point = FIXED_NODE_POINTS[row["id"]]
                    value = f"={lay.cell('structure', f'{point}_{f.key}')}"
            fill = styles.FORMULA_FILL if isinstance(value, str) and value.startswith("=") else None
            styles.put(ws, coord, value, fill or styles.INPUT_FILL)
    note = ws[f"A{t.end_row}"]
    note.value = (
        f"End of {t.name} ({t.capacity} rows). Keep the red guard row blank: anything typed "
        "there is counted as overflow on Checks. Spare rows: Active = 0, other cells blank."
    )
    note.font = styles.NOTE_FONT


def write_inputs(ws: Worksheet, lay: Layout, model: dict[str, Any], case: SceneCase) -> None:
    """Fill the Inputs sheet; rows beyond a table's capacity are counted, not written."""
    ws["A1"] = "Inputs -- the expanded crane model and the scene state"
    ws["A1"].font = styles.TITLE_FONT
    styles.legend(ws, 2)
    values = block_values(model, case)
    for name, block in lay.blocks.items():
        _write_block(ws, block, values[name])
    rows = table_rows(model)
    for name, table in lay.tables.items():
        _write_table(ws, lay, table, rows[name])
    ws.column_dimensions["A"].width = 34
    for col in "BCDEFGHIJKLMNOPQ":
        ws.column_dimensions[col].width = 16
