"""The OutOfPlane sheet: the app's out-of-plane (Y) pass, and its Results rows.

Source: ``calculation/engine.py``

* ``_out_of_plane_geometry`` -- the points: every document mass point (boom-local ones
  rotated about the heel, lateral offset ``v``), the lumped masses of
  ``crane_solver.lumped_point_masses`` (rope self-weight, idle hook blocks; y = 0) and every
  wind area (boom-local areas scaled by ``boom_wind_projection_scale`` = |sin(boom angle)|);
  the support-group origins boom_heel, front_base, rear_base, platform_heel.
* ``_out_of_plane_variants`` -- ``Fy = -m (ay' - roll' (z - z_ref))`` per mass, Y wind
  ``sign 0.5 rho Cd A_y V^2 / 1000``, the side force at the hoist's support node (else its
  hook node), the moments ``Mx = y Fz - (z - oz) Fy`` and ``Mz = (x - ox) Fy - y Fx`` about
  each group, the split ``0.5 total +/- M / width`` and the platform Fy and Mx.
* ``_augment_point_with_out_of_plane_loads`` -- the side force is ``sin(sidelead) x hook
  load``; the variants there are the linear decomposition of the single evaluation here.

Results rows: ``core/demand_extraction.py`` ``extract_connection_component_response``
(resultant, fradial, mradial, bearing line loads; the results table shows the minimum line
load as uplift ``max(0, -min)``, ``core/analysis_services.py``).

Every row formula is gated on ``active``; spare rows read 0 and text columns "".
"""

from __future__ import annotations

from collections.abc import Callable, Mapping
from functools import partial

from openpyxl.worksheet.worksheet import Worksheet

from .layout import Layout
from .layout_oop import OOP_GROUPS, OOP_NOTES, OOP_SPLITS
from .sheet_writers import write_block, write_table, write_title

_NO_DIAMETER = "n/a (no bearing diameter)"
#: Statics unknowns of the in-plane totals the split starts from.
_TOTALS = {
    "heel_point": ("h_x", "h_z"),
    "front_hinge": ("r_fb_x", "r_fb_z"),
    "rear_hinge": ("r_rb_x", "r_rb_z"),
}


def _scene_formulas(lay: Layout) -> dict[str, str]:
    geo = lambda key: lay.geo_cell("frame_motion", key)  # noqa: E731
    ld = lambda key: lay.loads_cell("load_scene", key)  # noqa: E731
    me = lambda key: lay.oop.cell("oop_scene", key)  # noqa: E731
    wind, speed = lay.cell("scene", "wind_direction"), lay.cell("scene", "wind_speed_m_s")
    width = lay.cell("structure", "boom_width_m")
    diameter = lay.cell("structure", "slewing_bearing_diameter_m")
    return {
        "ref_z_m": f"={ld('ref_z_m')}",
        "ay": f"={geo('solver_ay')}",
        "roll": f"={geo('roll_accel_crane')}",
        "wind_y_sign": f'=IF({speed}=0,0,IF({wind}="+Y",1,IF({wind}="-Y",-1,0)))',
        "wind_q_kN_m2": f"={ld('wind_q_kN_m2')}",
        "boom_wind_scale": f"={ld('boom_wind_scale')}",
        "hook_kN": f"={ld('hook_kN')}",
        "sidelead_kN": f"=SIN(RADIANS({lay.cell('scene', 'sidelead_deg')}))*{me('hook_kN')}",
        "boom_width_m": f"={width}+0",
        "width_ok": f"=IF({me('boom_width_m')}>0,1,0)",
        "bearing_d_m": f"={diameter}+0",
        "bearing_ok": f"=IF({me('bearing_d_m')}>0,1,0)",
    }


def _sidelead_formulas(lay: Layout) -> dict[str, str]:
    me = lambda key: lay.oop.cell("sidelead", key)  # noqa: E731
    hoists = lambda key: lay.col_range("hoists", key)  # noqa: E731
    pos = lambda key: lay.geo_col_range("node_positions", key)  # noqa: E731
    hoist = lay.cell("scene", "hoist_id")
    hoist_ids = lay.loads_col_range("hoist_ropes", "id_key")

    def node_row(id_key: str) -> str:
        return f'=IF({me(id_key)}="",0,IFERROR(MATCH({me(id_key)},{pos("id")},0),0))'

    def hoist_value(key: str) -> str:
        return f'=IF({me("hoist_row")}>0,INDEX({hoists(key)},{me("hoist_row")})&"","")'

    def at(col: str) -> str:
        return f"=IF({me('node_row')}>0,INDEX({pos(col)},{me('node_row')}),0)"

    return {
        "hoist_row": f'=IF({hoist}="",0,IFERROR(MATCH({hoist}&"",{hoist_ids},0),0))',
        "support_id": hoist_value("support_node_id"),
        "hook_id": hoist_value("load_node_id"),
        "support_row": node_row("support_id"),
        "hook_row": node_row("hook_id"),
        "node_row": f"=IF({me('support_row')}>0,{me('support_row')},{me('hook_row')})",
        "x_m": at("x_kin_m"),
        "z_m": at("z_kin_m"),
        "fy_kN": (f"=IF({me('node_row')}>0,{lay.oop.cell('oop_scene', 'sidelead_kN')},0)"),
    }


def _fy(lay: Layout, mass: str, z: str) -> str:
    """``-m (ay' - roll' (z - z_ref))`` (``_out_of_plane_variants`` ``fy_m``)."""
    s = lambda key: lay.oop.cell("oop_scene", key)  # noqa: E731
    return f"=-{mass}*({s('ay')}-{s('roll')}*({z}-{s('ref_z_m')}))"


def _mass_row(lay: Layout, i: int) -> dict[str, str]:
    inp = lambda key: lay.addr("mass_points", key, i)  # noqa: E731
    ld = lambda key: lay.loads_addr("mass_loads", key, i)  # noqa: E731
    me = lambda key: lay.oop.addr("oop_masses", key, i)  # noqa: E731
    on = f"{me('active')}=1"
    return {
        "active": f"={inp('active')}",
        "body": f"={ld('body')}",
        "frame": f'=IF({on},{inp("frame")}&"","")',
        "x_m": f"={ld('x_m')}",
        "y_m": f'=IF({on},IF({me("frame")}="boom_local",{inp("v_m")},{inp("y_m")})+0,0)',
        "z_m": f"={ld('z_m')}",
        "mass_t": f"={ld('mass_t')}",
        "fx_kN": f"={ld('fx')}",
        "fz_kN": f"={ld('fz')}",
        "fy_kN": _fy(lay, me("mass_t"), me("z_m")),
    }


def _lumped_row(lay: Layout, i: int) -> dict[str, str]:
    ld = lambda key: lay.loads_addr("node_loads", key, i)  # noqa: E731
    me = lambda key: lay.oop.addr("oop_lumped", key, i)  # noqa: E731
    return {
        "active": f"={lay.addr('nodes', 'active', i)}",
        "id": f"={ld('id')}",
        "x_m": f"={ld('x_m')}",
        "z_m": f"={ld('z_m')}",
        "rope_t": f"={ld('rope_mass_t')}",
        "idle_t": f"={ld('idle_mass_t')}",
        "mass_t": f"={me('rope_t')}+{me('idle_t')}",
        "fy_kN": _fy(lay, me("mass_t"), me("z_m")),
    }


def _wind_row(lay: Layout, i: int) -> dict[str, str]:
    inp = lambda key: lay.addr("wind_areas", key, i)  # noqa: E731
    ld = lambda key: lay.loads_addr("wind_loads", key, i)  # noqa: E731
    me = lambda key: lay.oop.addr("oop_winds", key, i)  # noqa: E731
    s = lambda key: lay.oop.cell("oop_scene", key)  # noqa: E731
    on = f"{me('active')}=1"
    return {
        "active": f"={inp('active')}",
        "body": f"={ld('body')}",
        "frame": f'=IF({on},{inp("frame")}&"","")',
        "x_m": f"={ld('x_m')}",
        "z_m": f"={ld('z_m')}",
        "cd": f"=IF({on},{inp('cd')}+0,0)",
        "area_y_m2": f"=IF({on},{inp('area_y_m2')}+0,0)",
        "scale": f'=IF({on},IF({me("frame")}="boom_local",{s("boom_wind_scale")},1),0)',
        "fy_kN": (
            f"={s('wind_y_sign')}*{s('wind_q_kN_m2')}*{me('cd')}*{me('area_y_m2')}*{me('scale')}"
        ),
    }


def _moment_row(lay: Layout, i: int) -> dict[str, str]:
    group, node = OOP_GROUPS[i - 1]
    me = lambda key: lay.oop.addr("moments", key, i)  # noqa: E731
    m = lambda key: lay.oop.col_range("oop_masses", key)  # noqa: E731
    lp = lambda key: lay.oop.col_range("oop_lumped", key)  # noqa: E731
    w = lambda key: lay.oop.col_range("oop_winds", key)  # noqa: E731
    side = lambda key: lay.oop.cell("sidelead", key)  # noqa: E731
    ox, oz = me("ox_m"), me("oz_m")
    row = {
        "group": group,
        "node": node,
        "ox_m": f"={lay.statics_cell('structure_nodes', f'{node}_x_m')}",
        "oz_m": f"={lay.statics_cell('structure_nodes', f'{node}_z_m')}",
        "mx_masses": (
            f"=SUMPRODUCT({m('y_m')},{m('fz_kN')})-SUMPRODUCT({m('z_m')}-{oz},{m('fy_kN')})"
        ),
        "mx_lumped": f"=-SUMPRODUCT({lp('z_m')}-{oz},{lp('fy_kN')})",
        "mx_wind": f"=-SUMPRODUCT({w('z_m')}-{oz},{w('fy_kN')})",
        "mx_side": f"=-({side('z_m')}-{oz})*{side('fy_kN')}",
        "mz_masses": (
            f"=SUMPRODUCT({m('x_m')}-{ox},{m('fy_kN')})-SUMPRODUCT({m('y_m')},{m('fx_kN')})"
        ),
        "mz_lumped": f"=SUMPRODUCT({lp('x_m')}-{ox},{lp('fy_kN')})",
        "mz_wind": f"=SUMPRODUCT({w('x_m')}-{ox},{w('fy_kN')})",
        "mz_side": f"=({side('x_m')}-{ox})*{side('fy_kN')}",
    }
    for comp in ("mx", "mz"):
        parts = ("masses", "lumped", "wind", "side")
        row[comp] = "=" + "+".join(me(f"{comp}_{part}") for part in parts)
    return row


def _split_row(lay: Layout, i: int) -> dict[str, str | int]:
    entity, group, total, side = OOP_SPLITS[i - 1]
    me = lambda key: lay.oop.addr("split", key, i)  # noqa: E731
    s = lambda key: lay.oop.cell("oop_scene", key)  # noqa: E731
    fx, fz = (lay.statics_value(key) for key in _TOTALS[total])
    ok, width = f"{s('width_ok')}=1", s("boom_width_m")
    return {
        "entity": entity,
        "group": group,
        "total_entity": total,
        "side": side,
        "total_fx": f"={fx}",
        "total_fz": f"={fz}",
        "mz": f"={lay.oop.moment(group, 'mz')}",
        "mx": f"={lay.oop.moment(group, 'mx')}",
        "fx": (f"=IF({ok},0.5*{me('total_fx')}+{me('side')}*{me('mz')}/{width},0)"),
        "fz": (f"=IF({ok},0.5*{me('total_fz')}+{me('side')}*{me('mx')}/{width},0)"),
    }


def _platform_formulas(lay: Layout) -> dict[str, str]:
    me = lambda key: lay.oop.cell("platform", key)  # noqa: E731
    parts = ("fy_masses", "fy_lumped", "fy_wind", "fy_side")
    return {
        "fy_masses": f"=SUM({lay.oop.col_range('oop_masses', 'fy_kN')})",
        "fy_lumped": f"=SUM({lay.oop.col_range('oop_lumped', 'fy_kN')})",
        "fy_wind": f"=SUM({lay.oop.col_range('oop_winds', 'fy_kN')})",
        "fy_side": f"={lay.oop.cell('sidelead', 'fy_kN')}",
        "fy_kN": "=" + "+".join(me(part) for part in parts),
        "mx_kNm": f"={lay.oop.moment('platform_base', 'mx')}",
    }


def oop_component_formula(lay: Layout, entity: str, component: str) -> str | None:
    """The number behind an out-of-plane ``reaction.<entity>.<component>`` (no leading
    ``=``), or ``None`` for an in-plane row."""
    if entity in ("heel_left", "heel_right"):
        fx, fz = lay.oop.split_value(entity, "fx"), lay.oop.split_value(entity, "fz")
        return {"fx": fx, "fz": fz, "resultant": f"SQRT({fx}^2+{fz}^2)"}[component]
    if entity != "platform_base" or component in ("fx", "fz", "my"):
        return None
    fy, mx = lay.oop.cell("platform", "fy_kN"), lay.oop.cell("platform", "mx_kNm")
    fx, fz, my = (lay.statics_value(key) for key in ("r_g_x", "r_g_z", "m_g"))
    mradial = f"SQRT({mx}^2+{my}^2)"
    d, ok = lay.oop.cell("oop_scene", "bearing_d_m"), lay.oop.cell("oop_scene", "bearing_ok")
    direct, ring = f"{fz}/(PI()*{d})", f"4*{mradial}/(PI()*{d}^2)"
    formulas = {
        "fy": fy,
        "mx": mx,
        "fradial": f"SQRT({fx}^2+{fy}^2)",
        "mradial": mradial,
        "max_bearing_line_load": f'IF({ok}=1,{direct}+{ring},"{_NO_DIAMETER}")',
        "min_bearing_line_load": f'IF({ok}=1,MAX(0,{ring}-{direct}),"{_NO_DIAMETER}")',
    }
    return formulas[component]


def oop_check_formulas(lay: Layout) -> dict[str, str]:
    """The ``zero_boom_width`` check."""
    return {"zero_boom_width": f"=IF({lay.oop.cell('oop_scene', 'width_ok')}=1,0,1)"}


_ROW_FORMULAS: dict[str, Callable[[Layout, int], Mapping[str, object]]] = {
    "oop_masses": _mass_row,
    "oop_lumped": _lumped_row,
    "oop_winds": _wind_row,
    "moments": _moment_row,
    "split": _split_row,
}


def write_oop(ws: Worksheet, lay: Layout) -> None:
    """Fill the OutOfPlane sheet (all formulas, driven by Inputs, Geometry, Loads, Statics)."""
    write_title(ws, "OutOfPlane -- Y loads, left/right split, platform Fy and Mx", OOP_NOTES)
    oop = lay.oop
    write_block(ws, oop.blocks["oop_scene"], _scene_formulas(lay))
    write_block(ws, oop.blocks["sidelead"], _sidelead_formulas(lay))
    for name, table in oop.tables.items():
        write_table(ws, table, partial(_ROW_FORMULAS[name], lay))
    write_block(ws, oop.blocks["platform"], _platform_formulas(lay))
    ws.column_dimensions["A"].width = 44
    for col in "BCDEFGHIJKLMN":
        ws.column_dimensions[col].width = 18
