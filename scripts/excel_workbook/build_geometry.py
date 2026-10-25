"""The Geometry sheet: global node positions and the crane-frame vessel motion.

Sources (the app code the formulas reproduce):

* node positions -- ``calculation/solver/kinematics.py`` ``KinematicsAdapter._compute_nodes``
  (boom-local nodes rotated about ``boom_heel``; global nodes as given; hanging nodes
  ``hanging_drop_m`` below their parent, a missing drop being 1 m) followed by
  ``calculation/rope_mass.py`` ``active_hook_positions`` / ``hook_positions_at_height``
  (the active hook's z = support z - max(0, tackle length), x unchanged). The support node
  is ``hook_support_node_id``: the hook node's ``below_node_id``, else the hoist's
  ``support_node_id``. ``core/scene_frame.py`` ``with_hook_tackle`` (what the Scene draws)
  gives the same positions.
* vessel motion -- ``calculation/vessel_frame.py`` ``crane_frame_motion``.

Every formula gates on ``active`` or on the frame text so spare rows give 0, never an
error.
"""

from __future__ import annotations

from openpyxl.worksheet.worksheet import Worksheet

from . import styles
from .layout import GEOMETRY_NOTES, Layout, Table
from .sheet_writers import local, write_block


def _scene_geometry_formulas(lay: Layout) -> dict[str, str]:
    scene = lambda key: lay.cell("scene", key)  # noqa: E731
    me = lambda key: lay.geo_cell("scene_geometry", key)  # noqa: E731
    ids = lay.geo_col_range("node_positions", "id")
    # Ids are compared as text (``&""``): a numeric hoist id typed on Inputs still matches.
    hoist_keys = lay.loads_col_range("hoist_ropes", "id_key")
    return {
        "boom_angle_rad": f"=RADIANS({scene('boom_angle_deg')})",
        "boom_cos": f"=COS({me('boom_angle_rad')})",
        "boom_sin": f"=SIN({me('boom_angle_rad')})",
        "heel_row": f'=IFERROR(MATCH("boom_heel",{ids},0),0)',
        "heel_x_m": f"=IF({me('heel_row')}>0,INDEX({lay.col_range('nodes', 'x_m')},"
        f"{me('heel_row')}),0)",
        "heel_z_m": f"=IF({me('heel_row')}>0,INDEX({lay.col_range('nodes', 'z_m')},"
        f"{me('heel_row')}),0)",
        "tackle_m": f"=MAX(0,{scene('tackle_length_m')})",
        "hook_id": (
            f'=IFERROR(INDEX({lay.col_range("hoists", "load_node_id")},'
            f'MATCH({scene("hoist_id")}&"",{hoist_keys},0))&"","")'
        ),
        "hook_row": f'=IF({me("hook_id")}="",0,IFERROR(MATCH({me("hook_id")},{ids},0),0))',
        "hook_parent_id": (
            f'=IF({me("hook_row")}>0,'
            f'INDEX({lay.col_range("nodes", "below_node_id")},{me("hook_row")})&"","")'
        ),
        "support_id": (
            f'=IF({me("hook_parent_id")}<>"",{me("hook_parent_id")},'
            f'IFERROR(INDEX({lay.col_range("hoists", "support_node_id")},'
            f'MATCH({scene("hoist_id")}&"",{hoist_keys},0))&"",""))'
        ),
        "support_row": (
            f'=IF({me("support_id")}="",0,IFERROR(MATCH({me("support_id")},{ids},0),0))'
        ),
        "support_z_m": (
            f"=IF({me('support_row')}>0,INDEX("
            f"{lay.geo_col_range('node_positions', 'z_kin_m')},{me('support_row')}),0)"
        ),
        "hook_moves": f"=IF(AND({me('hook_row')}>0,{me('support_row')}>0),1,0)",
    }


def _frame_motion_formulas(lay: Layout) -> dict[str, str]:
    scene = lambda key: lay.cell("scene", key)  # noqa: E731
    m = lambda key: lay.geo_cell("frame_motion", key)  # noqa: E731
    c, s, g = m("slew_cos"), m("slew_sin"), m("g_m_s2")
    return {
        "g_m_s2": f"=ABS({lay.cell('environment', 'gravity_m_s2')})",
        "roll_rad": f"=RADIANS({scene('vessel_roll_deg')})",
        "pitch_rad": f"=RADIANS({scene('vessel_pitch_deg')})",
        "slew_rad": f"=RADIANS({scene('crane_slewing_angle_deg')})",
        "slew_cos": f"=COS({m('slew_rad')})",
        "slew_sin": f"=SIN({m('slew_rad')})",
        "gx_vessel": f"={g}*SIN({m('pitch_rad')})",
        "gy_vessel": f"=-{g}*SIN({m('roll_rad')})*COS({m('pitch_rad')})",
        "gz_vessel": f"=-{g}*COS({m('roll_rad')})*COS({m('pitch_rad')})",
        "kx_vessel": f"={m('gx_vessel')}-{scene('accel_x_m_s2')}",
        "ky_vessel": f"={m('gy_vessel')}-{scene('accel_y_m_s2')}",
        "gx": f"={c}*{m('gx_vessel')}+{s}*{m('gy_vessel')}",
        "gy": f"=-{s}*{m('gx_vessel')}+{c}*{m('gy_vessel')}",
        "gz": f"={m('gz_vessel')}",
        "ax": f"={c}*{scene('accel_x_m_s2')}+{s}*{scene('accel_y_m_s2')}",
        "ay": f"=-{s}*{scene('accel_x_m_s2')}+{c}*{scene('accel_y_m_s2')}",
        "az": f"={scene('accel_z_m_s2')}",
        "kx": f"={m('gx')}-{m('ax')}",
        "ky": f"={m('gy')}-{m('ay')}",
        "kz": f"={m('gz')}-{m('az')}",
        "solver_ax": f"=-{m('kx')}",
        "solver_ay": f"=-{m('ky')}",
        "solver_az": f"=-{g}-{m('kz')}",
        "pitch_accel_crane": (
            f"=-{s}*{scene('roll_accel_rad_s2')}+{c}*{scene('pitch_accel_rad_s2')}"
        ),
        "roll_accel_crane": f"={c}*{scene('roll_accel_rad_s2')}+{s}*{scene('pitch_accel_rad_s2')}",
    }


def _node_row_formulas(lay: Layout, i: int) -> dict[str, str]:
    inp = lambda key: lay.addr("nodes", key, i)  # noqa: E731
    me = lambda key: lay.geo_addr("node_positions", key, i)  # noqa: E731
    sg = lambda key: lay.geo_cell("scene_geometry", key)  # noqa: E731
    ids = lay.geo_col_range("node_positions", "id")
    cos, sin = sg("boom_cos"), sg("boom_sin")
    is_local, hanging = f'{me("frame")}="boom_local"', f'{me("frame")}="hanging"'
    return {
        "active": f"={inp('active')}",
        "id": f'=IF({me("active")}=1,{inp("id")}&"","")',
        "frame": f'=IF({me("active")}=1,{inp("frame")}&"","")',
        "x_base_m": (
            f'=IF({is_local},{sg("heel_x_m")}+{inp("u_m")}*{cos}-{inp("w_m")}*{sin},'
            f'IF({me("frame")}="global",{inp("x_m")},0))'
        ),
        "z_base_m": (
            f'=IF({is_local},{sg("heel_z_m")}+{inp("u_m")}*{sin}+{inp("w_m")}*{cos},'
            f'IF({me("frame")}="global",{inp("z_m")},0))'
        ),
        "parent_row": (
            f'=IF(AND({hanging},{inp("below_node_id")}<>""),'
            f'IFERROR(MATCH({inp("below_node_id")}&"",{ids},0),0),0)'
        ),
        "drop_m": f'=IF({hanging},IF({inp("hanging_drop_m")}="",1,{inp("hanging_drop_m")}),0)',
        "x_kin_m": (
            f"=IF({hanging},IF({me('parent_row')}>0,"
            f"INDEX({lay.geo_col_range('node_positions', 'x_base_m')},{me('parent_row')}),0),"
            f"{me('x_base_m')})"
        ),
        "z_kin_m": (
            f"=IF({hanging},IF({me('parent_row')}>0,"
            f"INDEX({lay.geo_col_range('node_positions', 'z_base_m')},{me('parent_row')})"
            f"-{me('drop_m')},0),{me('z_base_m')})"
        ),
        "is_hook": (
            f"=IF(AND({sg('hook_moves')}=1,{me('active')}=1,{me('id')}={sg('hook_id')}),1,0)"
        ),
        "x_m": f"={me('x_kin_m')}",
        "z_m": f"=IF({me('is_hook')}=1,{sg('support_z_m')}-{sg('tackle_m')},{me('z_kin_m')})",
    }


def _write_table(ws: Worksheet, lay: Layout, table: Table) -> None:
    ws[f"A{table.title_row}"] = table.title
    ws[f"A{table.title_row}"].font = styles.SECTION_FONT
    for f in table.fields:
        col = table.column(f.key)
        header = ws[f"{col}{table.header_row}"]
        header.value = f"{f.header} [{f.unit}]" if f.unit else f.header
        header.font = styles.HEADER_FONT
        header.fill = styles.HEADER_FILL
        key = ws[f"{col}{table.key_row}"]
        key.value = f"{f.key} ({f.note})" if f.note else f.key
        key.font = styles.KEY_FONT
    for i in range(1, table.capacity + 1):
        for key, formula in _node_row_formulas(lay, i).items():
            styles.put(ws, local(lay.geo_addr(table.name, key, i)), formula)


def write_geometry(ws: Worksheet, lay: Layout) -> None:
    """Fill the Geometry sheet (all formulas, driven by Inputs)."""
    ws["A1"] = "Geometry -- node positions and crane-frame motion"
    ws["A1"].font = styles.TITLE_FONT
    for k, note in enumerate(GEOMETRY_NOTES):
        ws[f"A{3 + k}"] = note
        ws[f"A{3 + k}"].font = styles.NOTE_FONT
    write_block(ws, lay.geo_blocks["scene_geometry"], _scene_geometry_formulas(lay))
    write_block(ws, lay.geo_blocks["frame_motion"], _frame_motion_formulas(lay))
    _write_table(ws, lay, lay.geo_tables["node_positions"])
    ws.column_dimensions["A"].width = 44
    for col in "BCDEFGHIJKLMNOP":
        ws.column_dimensions[col].width = 18
