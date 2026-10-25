"""Recalculate an .xlsx workbook in pure Python using the ``formulas`` library."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

import formulas
import numpy as np
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries

Value = float | str | None

_REF = re.compile(r"^'?\[(?P<book>[^\]]+)\](?P<sheet>[^']+)'?!(?P<cell>[A-Z]+\d+(?::[A-Z]+\d+)?)$")


def _convert(raw: Any) -> Value:
    if raw is None:
        return None
    if isinstance(raw, (bool, np.bool_)):
        return float(raw)
    if isinstance(raw, (int, float, np.integer, np.floating)):
        return float(raw)
    if isinstance(raw, str):
        return None if raw == "" else raw
    return str(raw)


def _formula_cells(path: Path) -> dict[str, set[str]]:
    """Return {upper-case sheet name: {cell}} for every cell holding a formula."""
    wb = load_workbook(path)
    out: dict[str, set[str]] = {}
    for ws in wb.worksheets:
        cells = out.setdefault(ws.title.upper(), set())
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and value.startswith("="):
                    cells.add(cell.coordinate)
                elif hasattr(value, "ref") and hasattr(value, "text"):
                    for sub in ws[value.ref]:
                        cells.update(
                            c.coordinate for c in (sub if isinstance(sub, tuple) else (sub,))
                        )
    return out


def recalculate(path: Path) -> dict[str, dict[str, Value]]:
    """Evaluate every formula cell; return ``{sheet_name: {"A1": value}}``."""
    path = Path(path)
    wanted = _formula_cells(path)
    titles = {t.upper(): t for t in load_workbook(path).sheetnames}
    model = formulas.ExcelModel().loads(str(path)).finish()
    solution = model.calculate()

    result: dict[str, dict[str, Value]] = {title: {} for title in titles.values()}
    for key, rng in solution.items():
        match = _REF.match(str(key))
        if match is None:
            continue
        sheet = match["sheet"].upper()
        cell = match["cell"]
        values = np.atleast_2d(np.asarray(rng.value, dtype=object))
        min_col, min_row, _, _ = range_boundaries(cell)
        for (r, c), raw in np.ndenumerate(values):
            coord = f"{get_column_letter(min_col + c)}{min_row + r}"
            if coord in wanted.get(sheet, ()):
                result[titles[sheet]][coord] = _convert(raw)
    return result
