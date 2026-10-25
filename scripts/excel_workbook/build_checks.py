"""The Checks sheet: one live count per input check (0 = OK) and a total."""

from __future__ import annotations

from collections.abc import Mapping

from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from . import styles
from .layout import (
    CHECK_DESCRIPTIONS,
    CHECKS_HEADER_ROW,
    CHECKS_TOTAL_ROW,
    FIXED_NODE_POINTS,
    Layout,
    split,
)

_ID_TABLES = ("mass_points", "wind_areas", "nodes", "ropes", "hoists")
_BAD_FILL = PatternFill("solid", fgColor="F8CBAD")


def _active(lay: Layout, table: str) -> str:
    return f"({lay.col_range(table, 'active')}=1)"


def _missing(lay: Layout, target: str, refs: str) -> str:
    """Array term: 1 where ``refs`` names no active row of table ``target``."""
    return (
        f"(COUNTIFS({lay.col_range(target, 'id')},{refs},"
        f"{lay.col_range(target, 'active')},1)=0)"
    )


def _sp(*factors: str) -> str:
    return "SUMPRODUCT(" + "*".join(factors) + ")"


def input_check_formulas(
    lay: Layout, overflow_range: str, fixed_count_range: str
) -> dict[str, str]:
    """Formulas for the input checks of ``layout.CHECK_DESCRIPTIONS``."""
    r = lay.col_range
    duplicates = [
        _sp(
            _active(lay, t),
            f"({r(t, 'id')}<>\"\")",
            f"(COUNTIFS({r(t, 'id')},{r(t, 'id')},{r(t, 'active')},1)>1)",
        )
        for t in _ID_TABLES
    ]
    blanks = [_sp(_active(lay, t), f"({r(t, 'id')}=\"\")") for t in _ID_TABLES]
    blanks.append(_sp(_active(lay, "segments"), f"({r('segments', 'rope_id')}=\"\")"))
    node_refs = [
        _sp(_active(lay, "segments"), _missing(lay, "nodes", r("segments", "from_node_id"))),
        _sp(_active(lay, "segments"), _missing(lay, "nodes", r("segments", "to_node_id"))),
        _sp(_active(lay, "hoists"), _missing(lay, "nodes", r("hoists", "load_node_id"))),
        _sp(
            _active(lay, "hoists"),
            f"({r('hoists', 'support_node_id')}<>\"\")",
            _missing(lay, "nodes", r("hoists", "support_node_id")),
        ),
        _sp(
            _active(lay, "nodes"),
            f"({r('nodes', 'frame')}=\"hanging\")",
            _missing(lay, "nodes", r("nodes", "below_node_id")),
        ),
    ]
    rope_refs = [
        _sp(_active(lay, "segments"), _missing(lay, "ropes", r("segments", "rope_id"))),
        _sp(_active(lay, "hoists"), _missing(lay, "ropes", r("hoists", "rope_id"))),
    ]
    solve_count = f"COUNTIFS({r('ropes', 'active')},1,{r('ropes', 'role')},\"solve\")"
    hoist = lay.cell("scene", "hoist_id")
    wind = lay.cell("scene", "wind_direction")
    hoist_rope = f"INDEX({r('hoists', 'rope_id')},MATCH({hoist},{r('hoists', 'id')},0))"
    return {
        "overflow": f"=SUM({overflow_range})",
        "duplicate_ids": "=" + "+".join(duplicates),
        "blank_ids": "=" + "+".join(blanks),
        "node_references": "=" + "+".join(node_refs),
        "rope_references": "=" + "+".join(rope_refs),
        "hoist_without_segments": "="
        + _sp(
            _active(lay, "hoists"),
            f"(COUNTIFS({r('segments', 'rope_id')},{r('hoists', 'rope_id')},"
            f"{r('segments', 'active')},1)=0)",
        ),
        "multiple_solve_ropes": f"=MAX(0,{solve_count}-1)",
        "no_solve_rope": f"=IF({solve_count}=0,1,0)",
        "scene_hoist": (
            f"=IF(COUNTIFS({r('hoists', 'id')},{hoist},{r('hoists', 'active')},1)=0,1,"
            f"IF(COUNTIFS({r('ropes', 'id')},{hoist_rope},{r('ropes', 'active')},1,"
            f"{r('ropes', 'role')},\"fixed\")=0,1,0))"
        ),
        "wind_direction": (f'=IF(OR({wind}="+X",{wind}="-X",{wind}="+Y",{wind}="-Y"),0,1)'),
        "fixed_nodes_missing": f"=COUNTIF({fixed_count_range},0)",
    }


def _local(address: str) -> str:
    return split(address)[1]


def write_checks(
    ws: Worksheet, lay: Layout, extra_formulas: Mapping[str, str] | None = None
) -> None:
    """Write every check of ``layout.CHECK_DESCRIPTIONS``; later tasks pass their own formulas in
    ``extra_formulas`` (keyed like ``CHECK_DESCRIPTIONS``)."""
    ws["A1"] = "Checks -- every count must be 0"
    ws["A1"].font = styles.TITLE_FONT

    # Detail blocks below the check list: overflow per table, fixed nodes present.
    row = lay.checks_detail_row
    ws[f"A{row}"] = "Overflow per table"
    ws[f"A{row}"].font = styles.SECTION_FONT
    for col, text in zip("ABCD", ("Table", "Not written", "Typed in guard row", "Total")):
        ws[f"{col}{row + 1}"] = text
        ws[f"{col}{row + 1}"].font = styles.HEADER_FONT
    first = row + 2
    for i, table in enumerate(lay.tables.values()):
        rr = first + i
        ws[f"A{rr}"] = table.name
        styles.put(ws, f"B{rr}", f"={table.overflow_addr}")
        styles.put(ws, f"C{rr}", f"=COUNTA({table.guard_range})")
        styles.put(ws, f"D{rr}", f"=B{rr}+C{rr}")
    overflow_range = f"$D${first}:$D${first + len(lay.tables) - 1}"

    row = first + len(lay.tables) + 2
    ws[f"A{row}"] = "Fixed structure nodes (active rows with this id)"
    ws[f"A{row}"].font = styles.SECTION_FONT
    first = row + 1
    for i, node_id in enumerate(FIXED_NODE_POINTS):
        rr = first + i
        ws[f"A{rr}"] = node_id
        styles.put(
            ws,
            f"B{rr}",
            f"=COUNTIFS({lay.col_range('nodes', 'id')},A{rr},"
            f"{lay.col_range('nodes', 'active')},1)",
        )
    fixed_range = f"$B${first}:$B${first + len(FIXED_NODE_POINTS) - 1}"

    formulas = input_check_formulas(lay, overflow_range, fixed_range)
    formulas.update(extra_formulas or {})

    header = CHECKS_HEADER_ROW
    for col, text in zip("ABCD", ("Check", "Description", "Count", "Status")):
        ws[f"{col}{header}"] = text
        ws[f"{col}{header}"].font = styles.HEADER_FONT
        ws[f"{col}{header}"].fill = styles.HEADER_FILL
    for key, description in CHECK_DESCRIPTIONS.items():
        rr = lay.check_row(key)
        ws[f"A{rr}"] = key
        ws[f"B{rr}"] = description
        styles.put(ws, _local(lay.check(key)), formulas[key])
        styles.put(ws, f"D{rr}", f'=IF(C{rr}=0,"OK","CHECK")')
    last = lay.check_row(list(CHECK_DESCRIPTIONS)[-1])
    total = _local(lay.total_issues_addr)
    ws[f"B{CHECKS_TOTAL_ROW}"] = "Total issues"
    ws[f"B{CHECKS_TOTAL_ROW}"].font = styles.HEADER_FONT
    styles.put(ws, total, f"=SUM(C{header + 1}:C{last})")
    styles.put(ws, f"D{CHECKS_TOTAL_ROW}", f'=IF({total}=0,"OK","CHECK INPUTS")')
    bad = CellIsRule(operator="greaterThan", formula=["0"], fill=_BAD_FILL)
    ws.conditional_formatting.add(f"C{header + 1}:C{last}", bad)
    ws.conditional_formatting.add(total, bad)
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 90
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 16
