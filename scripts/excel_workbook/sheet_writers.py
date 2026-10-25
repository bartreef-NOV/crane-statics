"""Shared writers for formula sheets: labelled blocks, tables and the notes header.

Every formula-sheet writer uses ``local`` and ``write_block``; ``build_inputs`` keeps its own
block and table writers (they write typed values and validations, not formulas).
"""

from __future__ import annotations

from collections.abc import Callable, Mapping

from openpyxl.worksheet.worksheet import Worksheet

from . import styles
from .layout import Block, Table, split


def local(address: str) -> str:
    """``'Sheet!$B$5'`` -> ``'B5'``."""
    return split(address)[1]


def write_title(ws: Worksheet, title: str, notes: tuple[str, ...]) -> None:
    """Sheet title in A1, then one note per row from row 3."""
    ws["A1"] = title
    ws["A1"].font = styles.TITLE_FONT
    for k, note in enumerate(notes):
        ws[f"A{3 + k}"] = note
        ws[f"A{3 + k}"].font = styles.NOTE_FONT


def write_block(ws: Worksheet, block: Block, formulas: Mapping[str, str]) -> None:
    """A single-value block: label in A, formula in B, unit in C, key and note in D."""
    ws[f"A{block.title_row}"] = block.title
    ws[f"A{block.title_row}"].font = styles.SECTION_FONT
    for f in block.fields:
        row = block.row(f.key)
        ws[f"A{row}"] = f.header
        styles.put(ws, f"B{row}", formulas[f.key])
        ws[f"C{row}"] = f.unit or None
        ws[f"D{row}"] = f.key + (f"  -- {f.note}" if f.note else "")
        ws[f"D{row}"].font = styles.KEY_FONT


def write_table(
    ws: Worksheet, table: Table, row_values: Callable[[int], Mapping[str, object]]
) -> None:
    """A table: title, header with units, key row, then ``row_values(i)`` for every row."""
    ws[f"A{table.title_row}"] = table.title
    ws[f"A{table.title_row}"].font = styles.SECTION_FONT
    for f in table.fields:
        col = table.column(f.key)
        header = ws[f"{col}{table.header_row}"]
        header.value = f"{f.header} [{f.unit}]" if f.unit else f.header
        header.font = styles.HEADER_FONT
        header.fill = styles.HEADER_FILL
        key = ws[f"{col}{table.key_row}"]
        key.value = f"{f.key} ({f.note})" if f.note else f.key
        key.font = styles.KEY_FONT
    for i in range(1, table.capacity + 1):
        for key, value in row_values(i).items():
            styles.put(ws, local(table.addr(key, i)), value)
