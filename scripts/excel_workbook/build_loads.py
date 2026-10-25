"""The Loads sheet: in-plane loads on the four bodies, row by row, and per-body totals.

Sources (the app code the formulas reproduce; the solver wins where the Scene differs):

* mass points -- ``calculation/solver/crane_solver.py`` mass loop of ``_accumulate_loads``
  with ``inertia_force`` and ``rotate_boom_local_point``; split for reading into Gravity
  ``m g'`` and Inertia ``-m (a' + pitch' x lever)`` as ``core/drawn_loads.py``
  ``gravity``/``inertia`` draw them.
* wind -- ``_accumulate_loads`` ``_wind_fx``: ``sign 0.5 rho Cd A V^2 / 1000``, body
  ``boom`` areas times ``|sin(boom angle)|``; the sign is +1/-1 for +X/-X and 0 for Y wind
  (``calculation/engine.py`` ``split_wind_speed_components``).
* rope self-weight -- ``calculation/rope_mass.py`` ``rope_mass_at_nodes``, loaded at nodes
  like a mass point (``crane_solver.lumped_point_masses``).
* rope segment forces -- the rope loop of ``_accumulate_loads``: fixed ropes carry
  ``hook_kN / n_parts`` (scene hoist) or ``m_hook (|g| + az') / n_parts`` (idle hoists),
  the solve rope accumulates unit coefficients ``bh_*``.
* hook load, offlead and idle hook blocks -- the hook loop of ``_accumulate_loads`` and
  ``lumped_point_masses`` (``core/drawn_loads.py`` ``hook_load``/``idle_hook_mass``).

Every row formula is gated on ``active`` (or on ids found among active rows), so spare
rows read 0 and text columns read "".
"""

from __future__ import annotations

from collections.abc import Callable

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from . import styles
from .layout import BODIES, LOAD_CATEGORIES, LOADS_NOTES, Layout, Table
from .sheet_writers import local, write_block

Getter = Callable[[str], str]


def _scene_formulas(lay: Layout) -> dict[str, str]:
    geo = lambda key: lay.geo_cell("frame_motion", key)  # noqa: E731
    sg = lambda key: lay.geo_cell("scene_geometry", key)  # noqa: E731
    scene = lambda key: lay.cell("scene", key)  # noqa: E731
    me = lambda key: lay.loads_cell("load_scene", key)  # noqa: E731
    ids = lay.geo_col_range("node_positions", "id")
    wind = scene("wind_direction")
    return {
        "ref_x_m": f"={lay.cell('environment', 'reference_x_m')}+0",
        "ref_z_m": f"={lay.cell('environment', 'reference_z_m')}+0",
        "gx": f"={geo('gx')}",
        "gz": f"={geo('gz')}",
        "ax": f"={geo('ax')}",
        "az": f"={geo('az')}",
        "pitch": f"={geo('pitch_accel_crane')}",
        "idle_accel": f"={geo('g_m_s2')}+{geo('solver_az')}",
        "wind_sign": (
            f'=IF({scene("wind_speed_m_s")}=0,0,IF({wind}="+X",1,IF({wind}="-X",-1,'
            f'IF(OR({wind}="+Y",{wind}="-Y"),0,1))))'
        ),
        "wind_q_kN_m2": (
            f"=0.5*{lay.cell('environment', 'air_density_kg_m3')}"
            f"*{scene('wind_speed_m_s')}^2/1000"
        ),
        "boom_wind_scale": f"=ABS({sg('boom_sin')})",
        "hook_kN": f"={scene('hook_load_kN')}+0",
        "offlead_fx_kN": f"=SIN(RADIANS({scene('offlead_deg')}))*{me('hook_kN')}",
        "hook_row": f"={sg('hook_row')}",
        "anchor_row": (
            f'=IF({me("hook_row")}=0,0,IF({sg("hook_parent_id")}="",{me("hook_row")},'
            f'IFERROR(MATCH({sg("hook_parent_id")},{ids},0),{me("hook_row")})))'
        ),
        "scene_rope_id": (
            f'=IFERROR(INDEX({lay.loads_col_range("hoist_ropes", "rope_key")},'
            f'MATCH({scene("hoist_id")},{lay.loads_col_range("hoist_ropes", "id_key")},0))&"",'
            f'"")'
        ),
    }


def _mass_terms(m: str, x: str, z: str, lay: Layout) -> dict[str, str]:
    """Weight and inertia of mass ``m`` at (``x``, ``z``): drawn_loads gravity + inertia."""
    s = lambda key: lay.loads_cell("load_scene", key)  # noqa: E731
    dx, dz = f"({x}-{s('ref_x_m')})", f"({z}-{s('ref_z_m')})"
    return {
        "grav_fx": f"={m}*{s('gx')}",
        "grav_fz": f"={m}*{s('gz')}",
        "inertia_fx": f"=-{m}*({s('ax')}+{s('pitch')}*{dz})",
        "inertia_fz": f"=-{m}*({s('az')}-{s('pitch')}*{dx})",
    }


def _moment(x: str, z: str, fx: str, fz: str) -> str:
    return f"={x}*{fz}-{z}*{fx}"


def _mass_row(lay: Layout, i: int) -> dict[str, str]:
    inp = lambda key: lay.addr("mass_points", key, i)  # noqa: E731
    me = lambda key: lay.loads_addr("mass_loads", key, i)  # noqa: E731
    sg = lambda key: lay.geo_cell("scene_geometry", key)  # noqa: E731
    on, local = f"{me('active')}=1", f'{inp("frame")}="boom_local"'
    cos, sin = sg("boom_cos"), sg("boom_sin")
    row = {
        "active": f"={inp('active')}",
        "body": f'=IF({on},{inp("body")}&"","")',
        "x_m": (
            f"=IF({on},IF({local},{sg('heel_x_m')}+{inp('u_m')}*{cos}-{inp('w_m')}*{sin},"
            f"{inp('x_m')}+0),0)"
        ),
        "z_m": (
            f"=IF({on},IF({local},{sg('heel_z_m')}+{inp('u_m')}*{sin}+{inp('w_m')}*{cos},"
            f"{inp('z_m')}+0),0)"
        ),
        "mass_t": f"=IF({on},{inp('mass_t')}+0,0)",
    }
    row.update(_mass_terms(me("mass_t"), me("x_m"), me("z_m"), lay))
    x, z = me("x_m"), me("z_m")
    row["grav_m"] = _moment(x, z, me("grav_fx"), me("grav_fz"))
    row["inertia_m"] = _moment(x, z, me("inertia_fx"), me("inertia_fz"))
    row["fx"] = f"={me('grav_fx')}+{me('inertia_fx')}"
    row["fz"] = f"={me('grav_fz')}+{me('inertia_fz')}"
    row["m"] = f"={me('grav_m')}+{me('inertia_m')}"
    return row


def _wind_row(lay: Layout, i: int) -> dict[str, str]:
    inp = lambda key: lay.addr("wind_areas", key, i)  # noqa: E731
    me = lambda key: lay.loads_addr("wind_loads", key, i)  # noqa: E731
    sg = lambda key: lay.geo_cell("scene_geometry", key)  # noqa: E731
    s = lambda key: lay.loads_cell("load_scene", key)  # noqa: E731
    on, local = f"{me('active')}=1", f'{inp("frame")}="boom_local"'
    cos, sin = sg("boom_cos"), sg("boom_sin")
    return {
        "active": f"={inp('active')}",
        "body": f'=IF({on},{inp("body")}&"","")',
        "x_m": (
            f"=IF({on},IF({local},{sg('heel_x_m')}+{inp('u_m')}*{cos}-{inp('w_m')}*{sin},"
            f"{inp('x_m')}+0),0)"
        ),
        "z_m": (
            f"=IF({on},IF({local},{sg('heel_z_m')}+{inp('u_m')}*{sin}+{inp('w_m')}*{cos},"
            f"{inp('z_m')}+0),0)"
        ),
        "area_m2": f"=IF({on},{inp('area_x_m2')}+0,0)",
        "scale": f'=IF({on},IF({me("body")}="boom",{s("boom_wind_scale")},1),0)',
        "fx": (
            f"={s('wind_sign')}*{s('wind_q_kN_m2')}*IF({on},{inp('cd')}+0,0)"
            f"*{me('area_m2')}*{me('scale')}"
        ),
        "m": f"=-{me('z_m')}*{me('fx')}",
    }


def _rope_row(lay: Layout, i: int) -> dict[str, str]:
    inp = lambda key: lay.addr("ropes", key, i)  # noqa: E731
    me = lambda key: lay.loads_addr("rope_tensions", key, i)  # noqa: E731
    s = lambda key: lay.loads_cell("load_scene", key)  # noqa: E731
    seg = lambda key: lay.col_range("segments", key)  # noqa: E731
    on = f"{me('active')}=1"
    hook = me("hook_node_id")
    parts = (
        f"SUMIFS({seg('mult')},{seg('rope_id')},{me('id')},{seg('from_node_id')},{hook},"
        f"{seg('active')},1)+SUMIFS({seg('mult')},{seg('rope_id')},{me('id')},"
        f"{seg('to_node_id')},{hook},{seg('active')},1)"
    )
    return {
        "active": f"={inp('active')}",
        "id": f'=IF({on},{inp("id")}&"","")',
        "id_key": f"={me('id')}",
        "role": f'=IF({on},{inp("role")}&"","")',
        "is_solve": f'=IF(AND({on},{me("role")}="solve"),1,0)',
        "is_scene": (f'=IF(AND({on},{me("id")}<>"",{me("id")}={s("scene_rope_id")}),1,0)'),
        "hoist_row": (
            f'=IF(AND({on},{me("id")}<>""),IFERROR(MATCH({me("id")},'
            f'{lay.loads_col_range("hoist_ropes", "rope_key")},0),0),0)'
        ),
        "hook_node_id": (
            f'=IF({me("hoist_row")}>0,INDEX({lay.col_range("hoists", "load_node_id")},'
            f'{me("hoist_row")})&"","")'
        ),
        "n_parts": f'=IF({on},MAX(1,INT(IF({hook}="",0,{parts}))),0)',
        "hook_mass_t": (
            f"=IF({me('hoist_row')}>0,INDEX({lay.col_range('hoists', 'hook_mass_t')},"
            f"{me('hoist_row')})+0,0)"
        ),
        "tension_kN": (
            f"=IF(OR({me('active')}<>1,{me('is_solve')}=1),0,IF({me('is_scene')}=1,"
            f"{s('hook_kN')}/{me('n_parts')},IF({me('hoist_row')}>0,"
            f"{me('hook_mass_t')}*{s('idle_accel')}/{me('n_parts')},0)))"
        ),
    }


def _hoist_row(lay: Layout, i: int) -> dict[str, str]:
    inp = lambda key: lay.addr("hoists", key, i)  # noqa: E731
    me = lambda key: lay.loads_addr("hoist_ropes", key, i)  # noqa: E731
    sl = lambda key: lay.loads_col_range("segment_loads", key)  # noqa: E731
    on = f"{me('active')}=1"
    ids = lay.geo_col_range("node_positions", "id")
    return {
        "active": f"={inp('active')}",
        "id_key": f'=IF({on},{inp("id")}&"","")',
        "rope_key": f'=IF({on},{inp("rope_id")}&"","")',
        "is_scene": (
            f'=IF(AND({on},{me("id_key")}<>"",{me("id_key")}={lay.cell("scene", "hoist_id")}),'
            f"1,0)"
        ),
        "rope_ok": (
            f'=IF(AND({on},{me("rope_key")}<>""),IF(COUNTIFS('
            f'{lay.loads_col_range("rope_tensions", "id_key")},{me("rope_key")})>0,1,0),0)'
        ),
        "linear_mass_kg_m": (
            f"=IF(AND({me('rope_ok')}=1,{inp('linear_mass_kg_m')}>1E-9),"
            f"{inp('linear_mass_kg_m')}+0,0)"
        ),
        "hook_row": (
            f'=IF(AND({on},{inp("load_node_id")}<>""),'
            f'IFERROR(MATCH({inp("load_node_id")}&"",{ids},0),0),0)'
        ),
        "reeved_m": (
            f'=IF({me("rope_ok")}=1,SUMIFS({sl("reeved_m")},{sl("rope_id")},{me("rope_key")}),0)'
        ),
        "spool_node_id": (
            f'=IF({me("rope_ok")}=1,IFERROR(INDEX({sl("spool_end")},'
            f'MATCH({me("rope_key")},{sl("spool_key")},0))&"",""),"")'
        ),
        "spooled_m": (
            f'=IF(AND({me("linear_mass_kg_m")}>0,{me("spool_node_id")}<>"",'
            f'{inp("rope_length_m")}<>""),MAX(0,{inp("rope_length_m")}-{me("reeved_m")}),0)'
        ),
        "spooled_t": f"={me('spooled_m')}*{me('linear_mass_kg_m')}/1000",
        "idle_mass_t": (
            f"=IF(AND({on},{me('is_scene')}=0,{me('hook_row')}>0,{inp('hook_mass_t')}>0),"
            f"{inp('hook_mass_t')}+0,0)"
        ),
    }


def _segment_row(lay: Layout, i: int) -> dict[str, str]:
    inp = lambda key: lay.addr("segments", key, i)  # noqa: E731
    me = lambda key: lay.loads_addr("segment_loads", key, i)  # noqa: E731
    first = lambda key: lay.loads_addr("segment_loads", key, 1)  # noqa: E731
    rt = lambda key: lay.loads_col_range("rope_tensions", key)  # noqa: E731
    hr = lambda key: lay.loads_col_range("hoist_ropes", key)  # noqa: E731
    ids = lay.geo_col_range("node_positions", "id")
    x_col = lay.geo_col_range("node_positions", "x_m")
    z_col = lay.geo_col_range("node_positions", "z_m")
    spool = lay.col_range("nodes", "spool_node")
    on, ok = f"{me('active')}=1", f"{me('ok')}=1"

    def node_row(key: str) -> str:
        return f'=IF(AND({on},{inp(key)}<>""),IFERROR(MATCH({inp(key)}&"",{ids},0),0),0)'

    def at(col: str, row_key: str) -> str:
        return f"=IF({ok},INDEX({col},{me(row_key)}),0)"

    def flagged(row_key: str) -> str:
        return f"INDEX({spool},{me(row_key)})=1"

    def upto(key: str) -> str:
        """Expanding range from row 1 to this row (for running counts)."""
        return f"{first(key)}:{local(me(key))}"

    return {
        "active": f"={inp('active')}",
        "rope_id": f'=IF({on},{inp("rope_id")}&"","")',
        "role": (
            f'=IF(AND({on},{me("rope_id")}<>""),'
            f'IFERROR(INDEX({rt("role")},MATCH({me("rope_id")},{rt("id_key")},0))&"",""),"")'
        ),
        "from_row": node_row("from_node_id"),
        "to_row": node_row("to_node_id"),
        "ok": (f'=IF(AND({on},{me("role")}<>"",{me("from_row")}>0,{me("to_row")}>0),1,0)'),
        "xa_m": at(x_col, "from_row"),
        "za_m": at(z_col, "from_row"),
        "xb_m": at(x_col, "to_row"),
        "zb_m": at(z_col, "to_row"),
        "length_m": f"=SQRT(({me('xb_m')}-{me('xa_m')})^2+({me('zb_m')}-{me('za_m')})^2)",
        "ux": f"=IF({me('length_m')}<1E-9,0,({me('xb_m')}-{me('xa_m')})/{me('length_m')})",
        "uz": f"=IF({me('length_m')}<1E-9,0,({me('zb_m')}-{me('za_m')})/{me('length_m')})",
        "mult": f"=IF({ok},{inp('mult')}+0,0)",
        "tension_kN": (
            f'=IF(AND({ok},{me("role")}="fixed"),'
            f'INDEX({rt("tension_kN")},MATCH({me("rope_id")},{rt("id_key")},0)),0)'
        ),
        "fx": f"={me('tension_kN')}*{me('ux')}*{me('mult')}",
        "fz": f"={me('tension_kN')}*{me('uz')}*{me('mult')}",
        "bh_fx": f'=IF({me("role")}="solve",{me("ux")}*{me("mult")},0)',
        "bh_fz": f'=IF({me("role")}="solve",{me("uz")}*{me("mult")},0)',
        "reeved_m": f"=IF({ok},{me('length_m')}*MAX(1,INT({me('mult')})),0)",
        "linear_mass_kg_m": (
            f"=IF({ok},SUMIFS({hr('linear_mass_kg_m')},{hr('rope_key')},{me('rope_id')}),0)"
        ),
        "half_mass_t": f"={me('reeved_m')}*{me('linear_mass_kg_m')}/2000",
        "spool_end": (
            f'=IF({ok},IF({flagged("from_row")},{inp("from_node_id")}&"",'
            f'IF({flagged("to_row")},{inp("to_node_id")}&"","")),"")'
        ),
        "has_spool": f'=IF({me("spool_end")}<>"",1,0)',
        "spool_seq": (
            f"=IF({me('has_spool')}=1,SUMIFS({upto('has_spool')},{upto('rope_id')},"
            f"{me('rope_id')}),0)"
        ),
        "spool_key": f'=IF(AND({me("has_spool")}=1,{me("spool_seq")}=1),{me("rope_id")},"")',
    }


def _node_row(lay: Layout, i: int) -> dict[str, str]:
    inp = lambda key: lay.addr("nodes", key, i)  # noqa: E731
    geo = lambda key: lay.geo_addr("node_positions", key, i)  # noqa: E731
    me = lambda key: lay.loads_addr("node_loads", key, i)  # noqa: E731
    s = lambda key: lay.loads_cell("load_scene", key)  # noqa: E731
    sl = lambda key: lay.loads_col_range("segment_loads", key)  # noqa: E731
    hr = lambda key: lay.loads_col_range("hoist_ropes", key)  # noqa: E731
    seg = lambda key: lay.col_range("segments", key)  # noqa: E731
    on, node = f"{me('active')}=1", me("id")
    x, z = me("x_m"), me("z_m")

    def at_ends(col: str) -> str:
        """Sum of a segment column over segments starting here minus those ending here."""
        return (
            f"=IF({on},SUMIFS({sl(col)},{seg('from_node_id')},{node})"
            f"-SUMIFS({sl(col)},{seg('to_node_id')},{node}),0)"
        )

    row = {
        "active": f"={inp('active')}",
        "id": f"={geo('id')}",
        "body": f'=IF({on},{inp("body")}&"","")',
        "x_m": f"={geo('x_m')}",
        "z_m": f"={geo('z_m')}",
        "reeved_t": (
            f"=IF({on},SUMIFS({sl('half_mass_t')},{seg('from_node_id')},{node})"
            f"+SUMIFS({sl('half_mass_t')},{seg('to_node_id')},{node}),0)"
        ),
        "spooled_t": f"=IF({on},SUMIFS({hr('spooled_t')},{hr('spool_node_id')},{node}),0)",
        "rope_mass_t": f"={me('reeved_t')}+{me('spooled_t')}",
        "idle_mass_t": (
            f"=IF({on},SUMIFS({hr('idle_mass_t')},{lay.col_range('hoists', 'load_node_id')},"
            f"{node}),0)"
        ),
        "hook_fz": f"=IF(AND({on},{s('hook_row')}={i}),-{s('hook_kN')},0)",
        "offlead_fx": f"=IF(AND({on},{s('anchor_row')}={i}),{s('offlead_fx_kN')},0)",
        "rope_fx": at_ends("fx"),
        "rope_fz": at_ends("fz"),
        "bh_fx": at_ends("bh_fx"),
        "bh_fz": at_ends("bh_fz"),
    }
    rope = _mass_terms(me("rope_mass_t"), x, z, lay)
    row["rw_fx"] = f"{rope['grav_fx']}+({rope['inertia_fx'][1:]})"
    row["rw_fz"] = f"{rope['grav_fz']}+({rope['inertia_fz'][1:]})"
    idle = _mass_terms(me("idle_mass_t"), x, z, lay)
    row["idle_grav_fx"] = idle["grav_fx"]
    row["idle_grav_fz"] = idle["grav_fz"]
    row["idle_inertia_fx"] = idle["inertia_fx"]
    row["idle_inertia_fz"] = idle["inertia_fz"]
    for prefix, fx, fz in (
        ("rw", "rw_fx", "rw_fz"),
        ("idle_grav", "idle_grav_fx", "idle_grav_fz"),
        ("idle_inertia", "idle_inertia_fx", "idle_inertia_fz"),
        ("hook", "offlead_fx", "hook_fz"),
        ("rope", "rope_fx", "rope_fz"),
        ("bh", "bh_fx", "bh_fz"),
    ):
        row[f"{prefix}_m"] = _moment(x, z, me(fx), me(fz))
    known = ("rw", "idle_grav", "idle_inertia", "rope")
    row["fx"] = "=" + "+".join(me(f"{p}_fx") for p in known) + f"+{me('offlead_fx')}"
    row["fz"] = "=" + "+".join(me(f"{p}_fz") for p in known) + f"+{me('hook_fz')}"
    row["m"] = "=" + "+".join(me(f"{p}_m") for p in (*known, "hook"))
    return row


#: Body-total category -> [(Loads table, column prefix)] summed into it.
_CATEGORY_SOURCES: dict[str, tuple[tuple[str, str], ...]] = {
    "gravity": (("mass_loads", "grav"),),
    "rope_weight": (("node_loads", "rw"),),
    "idle_hook": (("node_loads", "idle_grav"),),
    "hook": (("node_loads", "hook"),),
    "wind": (("wind_loads", ""),),
    "inertia": (("mass_loads", "inertia"), ("node_loads", "idle_inertia")),
    "rope": (("node_loads", "rope"),),
}
#: Columns that hold a component, by (table, prefix, component); others count as 0.
_COMPONENT_COLUMN = {
    ("node_loads", "hook", "fx"): "offlead_fx",
    ("node_loads", "hook", "fz"): "hook_fz",
    ("wind_loads", "", "fz"): None,
}


def _body_row(lay: Layout, i: int) -> dict[str, str]:
    body = BODIES[i - 1]
    me = lambda key: lay.loads_addr("body_totals", key, i)  # noqa: E731

    def total(table: str, column: str) -> str:
        return (
            f"SUMIFS({lay.loads_col_range(table, column)},"
            f"{lay.loads_col_range(table, 'body')},{me('body')})"
        )

    row = {"body": body}
    for cat in LOAD_CATEGORIES:
        for comp in ("fx", "fz", "m"):
            terms = []
            for table, prefix in _CATEGORY_SOURCES[cat]:
                default = f"{prefix}_{comp}" if prefix else comp
                column = _COMPONENT_COLUMN.get((table, prefix, comp), default)
                if column is not None:
                    terms.append(total(table, column))
            row[f"{cat}_{comp}"] = "=" + ("+".join(terms) if terms else "0")
    for comp in ("fx", "fz", "m"):
        row[comp] = "=" + "+".join(me(f"{cat}_{comp}") for cat in LOAD_CATEGORIES)
        row[f"bh_{comp}"] = "=" + total("node_loads", f"bh_{comp}")
    return row


_ROW_FORMULAS: dict[str, Callable[[Layout, int], dict[str, str]]] = {
    "mass_loads": _mass_row,
    "wind_loads": _wind_row,
    "rope_tensions": _rope_row,
    "hoist_ropes": _hoist_row,
    "segment_loads": _segment_row,
    "node_loads": _node_row,
    "body_totals": _body_row,
}


def _write_table(ws: Worksheet, lay: Layout, table: Table) -> None:
    ws[f"A{table.title_row}"] = table.title
    ws[f"A{table.title_row}"].font = styles.SECTION_FONT
    for f in table.fields:
        col = table.column(f.key)
        header = ws[f"{col}{table.header_row}"]
        header.value = f"{f.header} [{f.unit}]" if f.unit else f.header
        header.font = styles.HEADER_FONT
        header.fill = styles.HEADER_FILL
        key = ws[f"{col}{table.key_row}"]
        key.value = f"{f.key} ({f.note})" if f.note else f.key
        key.font = styles.KEY_FONT
    rows = _ROW_FORMULAS[table.name]
    for i in range(1, table.capacity + 1):
        for key, formula in rows(lay, i).items():
            styles.put(ws, local(lay.loads_addr(table.name, key, i)), formula)


def write_loads(ws: Worksheet, lay: Layout) -> None:
    """Fill the Loads sheet (all formulas, driven by Inputs and Geometry)."""
    ws["A1"] = "Loads -- in-plane loads on the four bodies"
    ws["A1"].font = styles.TITLE_FONT
    for k, note in enumerate(LOADS_NOTES):
        ws[f"A{3 + k}"] = note
        ws[f"A{3 + k}"].font = styles.NOTE_FONT
    write_block(ws, lay.loads_blocks["load_scene"], _scene_formulas(lay))
    for table in lay.loads_tables.values():
        _write_table(ws, lay, table)
    ws.column_dimensions["A"].width = 44
    widest = max(len(t.fields) for t in lay.loads_tables.values())
    for index in range(2, widest + 1):
        ws.column_dimensions[get_column_letter(index)].width = 16
