"""Excel workbook: typos, numeric ids, edited row sheets and boom-angle extremes."""

import dataclasses
import math
from pathlib import Path
import re
from typing import Any, Callable

import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula

from scripts.excel_workbook import layout
from scripts.excel_workbook.build import build_workbook
from scripts.excel_workbook.reference import CASES

from tests.excel_workbook_helpers import (
    LOAD_TOL,
    TIGHT,
    _app_inplane,
    _assert_book_clean,
    _Book,
    _golden,
    _model,
    _workbook_results,
)


def _set_cell(path: Path, address: str, value: Any) -> None:
    wb = load_workbook(path)
    sheet, coord = layout.split(address)
    wb[sheet][coord] = value
    wb.save(path)


def _row(model: dict[str, Any], table: str, row_id: str) -> int:
    """1-based Inputs row of ``row_id`` in ``table`` (22689 model order)."""
    rows = {"nodes": model["nodes"], "ropes": model["ropes"]}.get(table)
    if rows is None:
        rows = [mp for b in model["bodies"] for mp in b["mass_points"]]
    return next(i for i, r in enumerate(rows, start=1) if r["id"] == row_id)


@pytest.mark.parametrize(
    ("table", "field", "row_id", "typo"),
    [
        ("mass_points", "body", "boom_structure", "Boom "),
        ("nodes", "frame", "mh_boom_sheave_3", "boom_locl"),
        ("ropes", "role", "SH__rope", "fixd"),
    ],
    ids=["body_trailing_space", "node_frame", "rope_role"],
)
def test_invalid_values_check_fires_on_typos(
    tmp_path: Path, table: str, field: str, row_id: str, typo: str
) -> None:
    model = _model()
    path = tmp_path / "typo.xlsx"
    build_workbook(model, CASES["base"], path, TIGHT)
    _set_cell(path, TIGHT.addr(table, field, _row(model, table, row_id)), typo)
    book = _Book(path)
    assert book(TIGHT.check("invalid_values")) == 1
    assert book(TIGHT.total_issues_addr) >= 1


@pytest.mark.parametrize("value", [2, "yes", "1"])
def test_invalid_values_check_fires_on_bad_active(tmp_path: Path, value: Any) -> None:
    path = tmp_path / "active.xlsx"
    build_workbook(_model(), CASES["base"], path, TIGHT)
    _set_cell(path, TIGHT.addr("mass_points", "active", 2), value)
    book = _Book(path)
    assert book(TIGHT.check("invalid_values")) == 1
    assert book(TIGHT.total_issues_addr) >= 1


def test_list_validations_reject_other_values(tmp_path: Path) -> None:
    path = tmp_path / "dv.xlsx"
    build_workbook(_model(), CASES["base"], path, TIGHT)
    validations = load_workbook(path)[layout.INPUTS].data_validations.dataValidation
    assert validations
    for dv in validations:
        assert dv.showErrorMessage and dv.errorStyle == "stop", dv.sqref
        # The one non-list validation: the bearing diameter must be a number > 0.
        assert dv.type == "list" or (dv.type, dv.formula1) == ("decimal", "0"), dv.sqref


@pytest.mark.parametrize(
    "mirror",
    [
        lambda lay: lay.loads_addr("mass_loads", "active", 1),
        lambda lay: lay.geo_addr("node_positions", "active", 3),
        lambda lay: lay.oop.addr("oop_winds", "active", 2),
        lambda lay: lay.results_addr("tensions", "active", 1),
    ],
    ids=["loads", "geometry", "out_of_plane", "results"],
)
def test_mirror_check_fires_when_a_row_sheet_is_out_of_step(
    tmp_path: Path, mirror: Callable[[layout.Layout], str]
) -> None:
    # An inserted or deleted Inputs row cannot be simulated with openpyxl (it does not
    # adjust formulas); a row sheet that no longer copies its Inputs rows is the same symptom.
    path = tmp_path / "mirror.xlsx"
    build_workbook(_model(), CASES["base"], path, TIGHT)
    assert _Book(path)(TIGHT.check("mirror")) == 0
    _set_cell(path, mirror(TIGHT), 0)
    assert _Book(path)(TIGHT.check("mirror")) == 1
    _set_cell(path, mirror(TIGHT), "=#REF!")  # a deleted Inputs row leaves #REF!
    assert _Book(path)(TIGHT.check("mirror")) == 1


_CELL_REF = re.compile(
    r"(?<![A-Za-z0-9_.!$])(?:(?P<sheet>[A-Za-z_][A-Za-z0-9_]*)!)?"
    r"(?P<c1>\$?[A-Z]{1,3})(?P<r1>\$?)(?P<n1>\d+)"
    r"(?::(?P<c2>\$?[A-Z]{1,3})(?P<r2>\$?)(?P<n2>\d+))?(?![A-Za-z0-9_(])"
)


def _shift_refs(formula: str, own_sheet: str, sheet: str, at: int) -> str:
    """``formula`` with every reference to rows >= ``at`` of ``sheet`` moved down one row."""

    def shift(m: re.Match[str]) -> str:
        if (m["sheet"] or own_sheet) != sheet:
            return m[0]
        out = (m["sheet"] + "!") if m["sheet"] else ""
        n1 = int(m["n1"])
        out += f"{m['c1']}{m['r1']}{n1 + 1 if n1 >= at else n1}"
        if m["n2"] is not None:
            n2 = int(m["n2"])
            out += f":{m['c2']}{m['r2']}{n2 + 1 if n2 >= at else n2}"
        return out

    return _CELL_REF.sub(shift, formula)


def _insert_row(path: Path, sheet: str, at: int) -> None:
    """Insert a blank row above sheet row ``at``, adjusting every formula as Excel does
    (openpyxl's ``insert_rows`` moves the cells but leaves the formulas alone)."""
    wb = load_workbook(path)
    wb[sheet].insert_rows(at)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and value.startswith("="):
                    cell.value = _shift_refs(value, ws.title, sheet, at)
                elif isinstance(value, ArrayFormula):
                    value.text = _shift_refs(value.text, ws.title, sheet, at)
                    if ws.title == sheet:
                        value.ref = _shift_refs(value.ref, sheet, sheet, at)
    wb.save(path)


def test_insert_row_helper_shifts_references() -> None:
    f = "=SUM(Inputs!$A$5:$A$9)+Inputs!B4+LOG10(C7)+Loads!$A$9+ATAN2(1,2)"
    assert _shift_refs(f, "Inputs", "Inputs", 5) == (
        "=SUM(Inputs!$A$6:$A$10)+Inputs!B4+LOG10(C8)+Loads!$A$9+ATAN2(1,2)"
    )


@pytest.mark.parametrize(
    ("table", "where"),
    [("nodes", "first"), ("nodes", "last"), ("mass_points", "first"), ("hoists", "last")],
)
def test_row_inserted_at_table_start_or_end_is_flagged(
    tmp_path: Path, table: str, where: str
) -> None:
    # Inserted as the first data row, or between the last data row and the guard row, the
    # row lies outside every table range: ROWS() is unchanged and the row sheets still copy
    # the shifted Inputs rows, so only the header-to-guard distance shows it.
    path = tmp_path / "inserted.xlsx"
    build_workbook(_model(), CASES["base"], path, TIGHT)
    t = TIGHT.tables[table]
    _insert_row(path, layout.INPUTS, t.first_row if where == "first" else t.guard_row)
    book = _Book(path)
    assert book(TIGHT.check("mirror")) == 1
    assert book(TIGHT.total_issues_addr) >= 1


def test_numeric_hoist_and_node_ids(tmp_path: Path) -> None:
    # MH renamed to the number 7, its support sheave to the number 4: the ids typed on
    # Inputs are numbers; every comparison must still match them.
    name = "outofplane_wind_roll_sidelead"
    model = _model()
    hoist = next(h for h in model["hoists"] if h["id"] == "MH")
    hoist["id"] = 7
    old = hoist["support_node_id"]
    for node in model["nodes"]:
        if node["id"] == old:
            node["id"] = 4
        if node["below_node_id"] == old:
            node["below_node_id"] = 4
    hoist["support_node_id"] = 4
    for rope in model["ropes"]:
        for seg in rope["segments"]:
            for end in ("from_node_id", "to_node_id"):
                if seg[end] == old:
                    seg[end] = 4
    case = dataclasses.replace(CASES[name], hoist_id=7)  # type: ignore[arg-type]
    path = tmp_path / "numeric.xlsx"
    build_workbook(model, case, path, TIGHT)
    book = _Book(path)
    _assert_book_clean(book, *layout.SHEETS)
    assert book(TIGHT.total_issues_addr) == 0
    got = _workbook_results(book, TIGHT, model)
    got["tension.MH"] = got.pop("tension.7")
    golden = _golden()["cases"][name]["app_results"]
    for key, value in got.items():
        assert math.isclose(value, golden[key], abs_tol=LOAD_TOL), (key, value, golden[key])


@pytest.mark.parametrize("angle", [0.0, 90.0])
def test_boom_angle_extremes(tmp_path: Path, angle: float) -> None:
    case = dataclasses.replace(CASES["inplane_wind_accel"], boom_angle_deg=angle)
    want = _app_inplane(case)
    assert want is not None
    path = tmp_path / "angle.xlsx"
    model = _model()
    build_workbook(model, case, path, TIGHT)
    book = _Book(path)
    _assert_book_clean(book, *layout.SHEETS)
    assert book(TIGHT.total_issues_addr) == 0
    got = _workbook_results(book, TIGHT, model)
    assert set(got) == set(want)
    for key, value in want.items():
        assert math.isclose(got[key], value, abs_tol=LOAD_TOL), (angle, key, got[key], value)
