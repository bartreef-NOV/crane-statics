"""Excel workbook acceptance (four 22689 cases, no error values) and delivery (CLI, file)."""

import dataclasses
import math
from pathlib import Path
from typing import Callable

from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula
import pytest

from scripts.excel_workbook import layout
from scripts.excel_workbook.reference import CASES, YAML_PATH

from tests.excel_workbook_helpers import (
    ACCEPT_TOL,
    CASE_NAMES,
    DELIVERED,
    EXPECTED_OMISSIONS,
    LAYOUT,
    _Book,
    _assert_book_clean,
    _build,
    _golden,
    _model,
    _same,
    _workbook_results,
)


@pytest.mark.parametrize("name", sorted(CASE_NAMES))
def test_all_four_cases_within_tolerance(case_book: Callable[[str], "_Book"], name: str) -> None:
    book = case_book(name)
    golden = _golden()["cases"][name]["app_results"]
    got = _workbook_results(book, LAYOUT, _model())
    assert set(golden) - set(got) == EXPECTED_OMISSIONS, name
    assert set(got) <= set(golden), set(got) - set(golden)
    for key, value in got.items():
        assert isinstance(value, float), (name, key, value)
        assert math.isclose(value, golden[key], rel_tol=0.0, abs_tol=ACCEPT_TOL), (
            name,
            key,
            value,
            golden[key],
        )
    assert book(LAYOUT.results_cell("results_status", "message")) == "OK"
    assert book(LAYOUT.total_issues_addr) == 0


@pytest.mark.parametrize("name", sorted(CASE_NAMES))
def test_no_error_values_anywhere(case_book: Callable[[str], "_Book"], name: str) -> None:
    # Only a singular matrix may show #NUM!, in the raw MMULT(MINVERSE) column on
    # Statics; none of the 22689 cases is singular, so every sheet must be error-free.
    book = case_book(name)
    assert set(book.values) == set(layout.SHEETS)
    _assert_book_clean(book, *layout.SHEETS)
    assert sum(len(v) for v in book.values.values()) > 1000


def test_workbook_recalculates_on_open(tmp_path: Path) -> None:
    assert load_workbook(_build(tmp_path)).calculation.fullCalcOnLoad is True


@pytest.mark.parametrize(("args", "case"), [((), "base"), (("--case", "idle_hoist"), "idle_hoist")])
def test_cli_writes_workbook_for_a_case(tmp_path: Path, args: tuple[str, ...], case: str) -> None:
    from scripts.excel_workbook.build import main

    assert next(iter(CASES)) == "base"
    out = tmp_path / "cli.xlsx"
    assert main([str(YAML_PATH), str(out), *args]) == 0
    wb = load_workbook(out)
    assert wb.sheetnames == list(layout.SHEETS)
    assert wb.calculation.fullCalcOnLoad is True
    for key, value in dataclasses.asdict(CASES[case]).items():
        sheet, coord = layout.split(LAYOUT.cell("scene", key))
        assert _same(wb[sheet][coord].value, value), (case, key)


def test_cli_rejects_unknown_case(tmp_path: Path) -> None:
    from scripts.excel_workbook.build import main

    with pytest.raises(SystemExit):
        main([str(YAML_PATH), str(tmp_path / "x.xlsx"), "--case", "no_such_case"])


def test_delivered_workbook_is_up_to_date(tmp_path: Path) -> None:
    # Regenerate with: python -m scripts.excel_workbook.build
    #   examples/crane_input_22689.yaml docs/crane-statics-validation-22689.xlsx
    fresh = load_workbook(_build(tmp_path))
    shipped = load_workbook(DELIVERED)
    assert shipped.sheetnames == fresh.sheetnames
    assert shipped.calculation.fullCalcOnLoad is True
    for ws in fresh.worksheets:
        other = shipped[ws.title]
        assert other.max_row == ws.max_row and other.max_column == ws.max_column, ws.title
        for row in ws.iter_rows():
            for cell in row:
                want, got = cell.value, other[cell.coordinate].value
                if isinstance(want, ArrayFormula):
                    assert isinstance(got, ArrayFormula), (ws.title, cell.coordinate)
                    assert (got.ref, got.text) == (want.ref, want.text), cell.coordinate
                else:
                    assert got == want, (ws.title, cell.coordinate, got, want)


def test_cli_capacity_option(tmp_path: Path) -> None:
    from scripts.excel_workbook.build import main

    out = tmp_path / "cap.xlsx"
    assert main([str(YAML_PATH), str(out), "--capacity", "nodes=31", "--capacity", "hoists=5"]) == 0
    lay = layout.Layout(layout.Capacities(nodes=31, hoists=5))
    wb = load_workbook(out)
    sheet, coord = layout.split(lay.tables["nodes"].capacity_addr)
    assert wb[sheet][coord].value == 31
    sheet, coord = layout.split(lay.addr("hoists", "active", 5))
    assert wb[sheet][coord].value == 0
    with pytest.raises(SystemExit):
        main([str(YAML_PATH), str(out), "--capacity", "no_table=3"])
