"""Excel workbook: the Checks sheet counts input problems."""

from pathlib import Path
from typing import Any

from openpyxl import load_workbook
import pytest

from scripts.excel_workbook import layout
from scripts.excel_workbook.build import build_workbook
from scripts.excel_workbook.reference import CASES

from tests.excel_workbook_helpers import LAYOUT, TIGHT, _Book, _build, _model, fit_layout


def test_overflow_check_trips(tmp_path: Path) -> None:
    model = _model()
    capacity = LAYOUT.tables["segments"].capacity
    extra = capacity + 1 - sum(len(r["segments"]) for r in model["ropes"])
    seg = model["ropes"][0]["segments"][0]
    model["ropes"][0]["segments"].extend(dict(seg) for _ in range(extra))
    path = _build(tmp_path, model)
    book = _Book(path)
    assert book(LAYOUT.tables["segments"].overflow_addr) == 1
    assert book(LAYOUT.check("overflow")) == 1
    assert book(LAYOUT.total_issues_addr) > 0
    assert book(LAYOUT.addr("segments", "active", capacity)) == 1


def test_overflow_check_counts_rows_typed_below_a_table(tmp_path: Path) -> None:
    path = _build(tmp_path)
    wb = load_workbook(path)
    sheet, coord = LAYOUT.tables["nodes"].guard_addr("id").replace("$", "").split("!")
    wb[sheet][coord] = "typed_below_the_table"
    wb.save(path)
    assert _Book(path)(LAYOUT.check("overflow")) == 1


def test_duplicate_id_check(tmp_path: Path) -> None:
    model = _model()
    model["nodes"].append(dict(model["nodes"][3]))
    assert _Book(_build(tmp_path, model))(LAYOUT.check("duplicate_ids")) > 0


def test_missing_node_reference_check(tmp_path: Path) -> None:
    model = _model()
    model["ropes"][0]["segments"][0]["to_node_id"] = "no_such_node"
    assert _Book(_build(tmp_path, model))(LAYOUT.check("node_references")) == 1


def test_inactive_node_reference_check(tmp_path: Path) -> None:
    path = _build(tmp_path)
    ids = [n["id"] for n in _model()["nodes"]]
    wb = load_workbook(path)
    row = ids.index("mh_aframe_sheave_1") + 1
    sheet, coord = LAYOUT.addr("nodes", "active", row).replace("$", "").split("!")
    wb[sheet][coord] = 0
    wb.save(path)
    # Two MH segments touch the deactivated sheave.
    assert _Book(path)(LAYOUT.check("node_references")) == 2


def test_two_solve_ropes_check(tmp_path: Path) -> None:
    model = _model()
    next(r for r in model["ropes"] if r["id"] == "MH__rope")["tension_mode"] = "solve"
    book = _Book(_build(tmp_path, model))
    assert book(LAYOUT.check("multiple_solve_ropes")) == 1
    assert book(LAYOUT.check("no_solve_rope")) == 0


def test_hoist_rope_without_active_segments_check(tmp_path: Path) -> None:
    model = _model()
    next(r for r in model["ropes"] if r["id"] == "SH__rope")["segments"] = []
    assert _Book(_build(tmp_path, model))(LAYOUT.check("hoist_without_segments")) == 1


def test_hanging_chain_check(tmp_path: Path) -> None:
    model = _model()
    parent = next(n for n in model["nodes"] if n["frame"] == "hanging")
    child = dict(parent, id="chained_hook", below_node_id=parent["id"])
    model["nodes"].append(child)
    assert _Book(_build(tmp_path, model))(LAYOUT.check("hanging_chain")) == 1


def test_duplicate_label_ids_are_allowed(tmp_path: Path) -> None:
    # Mass-point and wind-area ids are labels only; the app accepts duplicates there.
    model = _model()
    for body in model["bodies"]:
        for key in ("mass_points", "wind_areas"):
            if len(body[key]) > 1:
                body[key][1]["id"] = body[key][0]["id"]
    path = tmp_path / "labels.xlsx"
    build_workbook(model, CASES["base"], path, TIGHT)
    book = _Book(path)
    assert book(TIGHT.check("duplicate_ids")) == 0
    assert book(TIGHT.total_issues_addr) == 0


def _set_cells(path: Path, values: dict[str, Any]) -> None:
    wb = load_workbook(path)
    for address, value in values.items():
        sheet, coord = address.replace("$", "").split("!")
        wb[sheet][coord] = value
    wb.save(path)


@pytest.mark.parametrize("diameter", [0, -1])
def test_bearing_diameter_must_be_positive(tmp_path: Path, diameter: float) -> None:
    # The app rejects a slewing bearing diameter <= 0 when it validates the document.
    path = tmp_path / "bearing.xlsx"
    build_workbook(_model(), CASES["base"], path, TIGHT)
    _set_cells(path, {TIGHT.cell("structure", "slewing_bearing_diameter_m"): diameter})
    book = _Book(path)
    assert book(TIGHT.check("invalid_bearing_diameter")) == 1
    assert book(TIGHT.total_issues_addr) >= 1


def test_bearing_diameter_cell_rejects_non_positive_entries(tmp_path: Path) -> None:
    path = tmp_path / "bearing_dv.xlsx"
    build_workbook(_model(), CASES["base"], path, TIGHT)
    coord = TIGHT.cell("structure", "slewing_bearing_diameter_m").replace("$", "").split("!")[1]
    validations = load_workbook(path)[layout.INPUTS].data_validations.dataValidation
    dv = next(v for v in validations if coord in v.sqref)
    assert (dv.type, dv.operator, dv.formula1) == ("decimal", "greaterThan", "0")
    assert dv.allow_blank and dv.showErrorMessage and dv.errorStyle == "stop"


def test_boom_local_on_non_boom_body_is_flagged(tmp_path: Path) -> None:
    # The app ignores u/v/w on a body other than the boom (it uses x/z); the workbook would
    # rotate the point with the boom, so the combination is flagged on every point table.
    model = _model()
    lay = fit_layout(model, spare=1)
    path = tmp_path / "local.xlsx"
    build_workbook(model, CASES["base"], path, lay)
    masses = [mp for b in model["bodies"] for mp in b["mass_points"]]
    mass_row = next(i for i, mp in enumerate(masses, 1) if mp["frame"] == "boom_local")
    winds = [wa for b in model["bodies"] for wa in b["wind_areas"]]
    wind_row = next(i for i, wa in enumerate(winds, 1) if wa["frame"] == "boom_local")
    node_row = next(i for i, n in enumerate(model["nodes"], 1) if n["frame"] == "boom_local")
    spare = len(model["nodes"]) + 1
    _set_cells(
        path,
        {
            lay.addr("mass_points", "body", mass_row): "rear_leg",
            lay.addr("wind_areas", "body", wind_row): "front_leg",
            lay.addr("nodes", "body", node_row): "platform",
            # A spare row (Active = 0) is never counted.
            lay.addr("nodes", "body", spare): "front_leg",
            lay.addr("nodes", "frame", spare): "boom_local",
        },
    )
    book = _Book(path)
    assert book(lay.check("boom_local_non_boom")) == 3
    assert book(lay.total_issues_addr) >= 3
