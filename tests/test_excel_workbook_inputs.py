"""Excel workbook: the Inputs tables, sheet order, capacities and the bearing-diameter check."""

import dataclasses
from pathlib import Path
from typing import Callable

from openpyxl import load_workbook
import pytest

from scripts.excel_workbook import layout
from scripts.excel_workbook.build import build_workbook
from scripts.excel_workbook.reference import CASES

from tests.excel_workbook_helpers import (
    CASE_NAMES,
    DERIVED_HOIST_KEYS,
    LAYOUT,
    OMITTED_ROPE_KEYS,
    TIGHT,
    _Book,
    _build,
    _expected_rows,
    _model,
    _same,
)


def test_inputs_round_trip(case_book: Callable[[str], "_Book"]) -> None:
    model = _model()
    book = case_book("base")
    rows = _expected_rows(model)
    assert set(rows) == set(LAYOUT.tables)
    assert set(model["hoists"][0]) - set(LAYOUT.tables["hoists"].keys) == DERIVED_HOIST_KEYS
    assert set(model["ropes"][0]) - set(LAYOUT.tables["ropes"].keys) == OMITTED_ROPE_KEYS
    for name, table in LAYOUT.tables.items():
        assert len(rows[name]) <= table.capacity, name
        for i, row in enumerate(rows[name], start=1):
            assert book(LAYOUT.addr(name, "active", i)) == 1, (name, i)
            for key in table.keys:
                if key != "active":
                    got = book(LAYOUT.addr(name, key, i))
                    assert _same(got, row[key]), (name, i, key, got, row[key])
        for i in range(len(rows[name]) + 1, table.capacity + 1):
            assert book(LAYOUT.addr(name, "active", i)) == 0, (name, i)
            for key in table.keys:
                if key != "active":
                    assert book(LAYOUT.addr(name, key, i)) is None, (name, i, key)
        assert book(table.overflow_addr) == 0, name

    env = model["environment"]
    for key in ("gravity_m_s2", "air_density_kg_m3", "max_wind_speed_m_s"):
        assert _same(book(LAYOUT.cell("environment", key)), env[key]), key
    for key in ("reference_x_m", "reference_z_m"):
        assert _same(book(LAYOUT.cell("environment", key)), env["vessel_motion"][key]), key
    structure = model["structure"]
    for key in ("boom_width_m", "slewing_bearing_diameter_m"):
        assert _same(book(LAYOUT.cell("structure", key)), structure[key]), key
    for point, values in structure["points"].items():
        for axis in ("x_m", "z_m"):
            got = book(LAYOUT.cell("structure", f"{point}_{axis}"))
            assert _same(got, values[axis]), (point, axis)


@pytest.mark.parametrize("name", sorted(CASE_NAMES))
def test_scene_state_cells_hold_the_case(case_book: Callable[[str], "_Book"], name: str) -> None:
    book = case_book(name)
    for key, value in dataclasses.asdict(CASES[name]).items():
        assert _same(book(LAYOUT.cell("scene", key)), value), (name, key)


def test_checks_are_clean_for_22689(case_book: Callable[[str], "_Book"]) -> None:
    book = case_book("base")
    for key in layout.CHECK_DESCRIPTIONS:
        assert book(LAYOUT.check(key)) == 0, key
    assert book(LAYOUT.total_issues_addr) == 0


def test_custom_capacities(tmp_path: Path) -> None:
    lay = layout.Layout(layout.Capacities(nodes=30, segments=16, hoists=5))
    assert lay.tables["segments"].first_row != LAYOUT.tables["segments"].first_row
    path = tmp_path / "small.xlsx"
    build_workbook(_model(), CASES["base"], path, lay)
    book = _Book(path)
    assert book(lay.addr("segments", "active", 16)) == 1
    assert book(lay.addr("nodes", "active", 30)) == 0
    assert book(lay.total_issues_addr) == 0


def test_workbook_sheets_in_order(tmp_path: Path) -> None:
    wb = load_workbook(_build(tmp_path))
    assert wb.sheetnames == list(layout.SHEETS)
    assert layout.SHEETS == (
        "Conventions",
        "Inputs",
        "Geometry",
        "Loads",
        "Statics",
        "OutOfPlane",
        "Results",
        "Checks",
    )


@pytest.mark.parametrize(
    ("diameter", "count"), [(0.0, 1), (-2.0, 1), ("abc", 1), (None, 0), (12.5, 0)]
)
def test_invalid_bearing_diameter_check(
    tmp_path: Path, diameter: float | str | None, count: int
) -> None:
    model = _model()
    model["structure"]["slewing_bearing_diameter_m"] = diameter
    path = tmp_path / "bearing.xlsx"
    build_workbook(model, CASES["base"], path, TIGHT)
    book = _Book(path)
    assert book(TIGHT.check("invalid_bearing_diameter")) == count
    assert book(TIGHT.total_issues_addr) == count
