"""Shared helpers for the Excel validation workbook tests (tests/test_excel_workbook_*.py)."""

import copy
import dataclasses
import json
import math
from pathlib import Path
from typing import Any, Callable

from openpyxl import load_workbook
import pytest

from crane_statics.calculation.engine import split_wind_speed_components
from crane_statics.calculation.rope_mass import (
    active_hook_positions,
    falls_override_effective_mult,
    rope_mass_at_nodes,
)
from crane_statics.calculation.solver import crane_solver
from crane_statics.calculation.solver.kinematics import KinematicsAdapter
from crane_statics.calculation.vessel_frame import VesselInputs, crane_frame_motion
from crane_statics.core import drawn_loads
from crane_statics.core.io import load_document
from scripts.excel_workbook import layout
from scripts.excel_workbook.build import build_workbook
from scripts.excel_workbook.recalc import recalculate
from scripts.excel_workbook.reference import (
    CASES,
    GOLDEN_PATH,
    SceneCase,
    YAML_PATH,
    _scene_inputs,
    _solve,
    app_results,
    expanded_model,
)

CASE_NAMES = {"base", "inplane_wind_accel", "outofplane_wind_roll_sidelead", "idle_hoist"}


def _num(value: float | str | None) -> float:
    assert isinstance(value, float), value
    return value


def _golden() -> dict[str, Any]:
    golden: dict[str, Any] = json.loads(GOLDEN_PATH.read_text(encoding="utf-8"))
    return golden


LAYOUT = layout.DEFAULT_LAYOUT


#: expanded_model keys deliberately not typed on Inputs (derived in later sheets, or
#: kept once on the hoist row -- the app reads the hoist's linear mass).
DERIVED_HOIST_KEYS = {"is_payload", "parts_at_hook", "spool_node_id"}


OMITTED_ROPE_KEYS = {"tension_mode", "linear_mass_kg_m", "segments"}


def _model() -> dict[str, Any]:
    return copy.deepcopy(expanded_model(YAML_PATH))


def _build(tmp_path: Path, model: dict[str, Any] | None = None, case: str = "base") -> Path:
    path = tmp_path / f"wb_{case}.xlsx"
    build_workbook(_model() if model is None else model, CASES[case], path)
    return path


class _Book:
    """Read cells by sheet-qualified address; formula cells are recalculated."""

    def __init__(self, path: Path) -> None:
        self.wb = load_workbook(path)
        self.values = recalculate(path)

    def __call__(self, address: str) -> Any:
        sheet, coord = address.replace("$", "").split("!")
        raw = self.wb[sheet][coord].value
        if isinstance(raw, str) and raw.startswith("="):
            return self.values[sheet][coord]
        return raw


def _expected_rows(model: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    return {
        "mass_points": [
            {"body": b["id"], **mp} for b in model["bodies"] for mp in b["mass_points"]
        ],
        "wind_areas": [{"body": b["id"], **wa} for b in model["bodies"] for wa in b["wind_areas"]],
        "nodes": model["nodes"],
        "segments": [{"rope_id": r["id"], **s} for r in model["ropes"] for s in r["segments"]],
        "ropes": [{"id": r["id"], "role": r["tension_mode"]} for r in model["ropes"]],
        "hoists": model["hoists"],
    }


def _same(got: Any, expected: Any) -> bool:
    if expected is None:
        return got is None
    if isinstance(expected, bool):
        return bool(got == (1 if expected else 0))
    if isinstance(expected, (int, float)):
        return isinstance(got, (int, float)) and math.isclose(got, expected, abs_tol=1e-12)
    return bool(got == expected)


def _solver_loads(
    case: SceneCase, mutate: Callable[[dict[str, Any]], None] | None = None
) -> tuple[dict[str, dict], dict[str, dict], dict[str, Any]]:
    """The solver's own per-body in-plane loads ``(ld, bh)`` at ``case``.

    ``_accumulate_loads`` evaluated directly with the inputs ``_angle_invariants`` and
    ``compute_analysis_point`` give it (the app composes the same linear terms from unit
    channels). Returns the positions and the document too.
    """
    doc = load_document(YAML_PATH).model_dump()
    if mutate is not None:
        mutate(doc)
    positions = active_hook_positions(
        doc,
        case.hoist_id,
        KinematicsAdapter(doc).compute_all_node_positions(case.boom_angle_deg),
        case.tackle_length_m,
    )
    motion = crane_frame_motion(
        abs(float(doc["environment"]["gravity_m_s2"])),
        VesselInputs(
            accel_x_m_s2=case.accel_x_m_s2,
            accel_y_m_s2=case.accel_y_m_s2,
            accel_z_m_s2=case.accel_z_m_s2,
            roll_deg=case.vessel_roll_deg,
            pitch_deg=case.vessel_pitch_deg,
            slew_deg=case.crane_slewing_angle_deg,
            roll_accel_rad_s2=case.roll_accel_rad_s2,
            pitch_accel_rad_s2=case.pitch_accel_rad_s2,
        ),
    )
    rope_mass = {
        nid: entry.mass_t
        for nid, entry in rope_mass_at_nodes(
            doc, positions, effective_mult=falls_override_effective_mult(doc, {})
        ).items()
    }
    wind_x, _wind_y = split_wind_speed_components(case.wind_speed_m_s, case.wind_direction)
    ld, bh = crane_solver._accumulate_loads(
        doc,
        {nid: (float(p["x"]), float(p["z"])) for nid, p in positions.items()},
        {n["id"]: n.get("body", "") for n in doc["nodes"]},
        case.hoist_id,
        case.hook_load_kN,
        wind_x,
        motion.accel_x_m_s2,
        motion.accel_z_m_s2,
        motion.pitch_accel_rad_s2,
        offlead_lateral_force_kN=math.sin(math.radians(case.offlead_deg)) * case.hook_load_kN,
        solver_ctx=crane_solver.build_solver_static_context(doc),
        rope_mass_t_by_node=rope_mass,
        boom_angle_deg=case.boom_angle_deg,
    )
    return ld, bh, {"doc": doc, "positions": positions}


def _drawn_category_sums(case: SceneCase) -> tuple[dict[str, dict[str, dict]], float]:
    """{body: {category: {"fx", "fz", "m"}}} of the Scene's free-body loads, and T_bh."""
    session, calculator, (pt, _rows, _conn) = _solve(YAML_PATH, case)
    inputs = _scene_inputs(session, calculator, pt)
    assert inputs.document_dict is not None
    positions = inputs.node_positions()
    out: dict[str, dict[str, dict]] = {}
    for body in layout.BODIES:
        loads = drawn_loads.build_body_loads(
            inputs.document_dict,
            hoist=dict(inputs.hoist),
            conditions=inputs.conditions,
            node_positions=positions,
            analysis_point=pt,
            body_filter=drawn_loads.FreeBodyFilter(
                body, drawn_loads.node_body_map(inputs.document_dict)
            ),
        )
        out[body] = {
            cat: {
                "fx": sum(v.fx_kN for v in loads.category(title)),
                "fz": sum(v.fz_kN for v in loads.category(title)),
                "m": sum(v.world_x * v.fz_kN - v.world_z * v.fx_kN for v in loads.category(title)),
            }
            for cat, title in layout.LOAD_CATEGORIES.items()
        }
    t_bh = app_results(YAML_PATH, case)["tension.BH"]
    return out, t_bh


LOAD_TOL = 0.01


def _assert_body_totals_match_solver(
    book: "_Book", lay: layout.Layout, ld: dict[str, dict], bh: dict[str, dict], tag: Any
) -> None:
    for body in layout.BODIES:
        for comp, key in (("Fx", "fx"), ("Fz", "fz"), ("M", "m")):
            got = book(lay.body_total(body, key))
            assert math.isclose(got, ld[body][comp], abs_tol=LOAD_TOL), (tag, body, key)
            got = book(lay.body_total(body, f"bh_{key}"))
            assert math.isclose(got, bh[body][comp], abs_tol=1e-9), (tag, body, f"bh_{key}")


def _assert_no_errors(source: "Path | _Book", sheet: str = layout.LOADS) -> dict[str, Any]:
    """No error values on ``sheet``; reuses a ``_Book``'s values instead of recalculating."""
    values = source.values[sheet] if isinstance(source, _Book) else recalculate(source)[sheet]
    errors = {c: v for c, v in values.items() if isinstance(v, str) and v.startswith("#")}
    assert not errors, list(errors.items())[:10]
    return values


def _rename_node(model: dict[str, Any], old: str, new: str) -> None:
    for node in model["nodes"]:
        if node["id"] == old:
            node["id"] = new
    for rope in model["ropes"]:
        for seg in rope["segments"]:
            for end in ("from_node_id", "to_node_id"):
                if seg[end] == old:
                    seg[end] = new


def _wind_case(direction: str, speed: float = 15.1) -> SceneCase:
    return dataclasses.replace(
        CASES["inplane_wind_accel"], wind_direction=direction, wind_speed_m_s=speed
    )


#: Tight capacities (22689 fits exactly) keep the per-direction rebuilds quick.
TIGHT = layout.Layout(
    layout.Capacities(mass_points=9, wind_areas=8, nodes=29, segments=16, ropes=4, hoists=4)
)


def _build_case(tmp_path: Path, case: SceneCase, lay: layout.Layout = TIGHT) -> Path:
    path = tmp_path / f"wb_{case.wind_direction.replace('+', 'p').replace('-', 'm')}.xlsx"
    build_workbook(_model(), case, path, lay)
    return path


#: Golden reaction components the workbook must reproduce in-plane. The heel point's
#: boom_axial/normal (name-based axis node) and the out-of-plane rows are not in-plane.
INPLANE_REQUIRED = {
    ("heel_point", "fx"),
    ("heel_point", "fz"),
    ("rear_hinge", "fx"),
    ("rear_hinge", "fz"),
    ("front_hinge", "fx"),
    ("front_hinge", "fz"),
    ("upper_hinge", "fx"),
    ("upper_hinge", "fz"),
    ("platform_base", "fx"),
    ("platform_base", "fz"),
    ("platform_base", "my"),
}


def _workbook_results(book: "_Book", lay: layout.Layout, model: dict[str, Any]) -> dict[str, Any]:
    """The Results sheet as ``app_results`` keys (``tension.<hoist>``, ``reaction.*``)."""
    out: dict[str, Any] = {}
    for i, hoist in enumerate(model["hoists"], start=1):
        assert book(lay.results_addr("tensions", "hoist_id", i)) == hoist["id"]
        out[f"tension.{hoist['id']}"] = book(lay.results_addr("tensions", "tension_kN", i))
    for entity, component in lay.result_rows:
        out[f"reaction.{entity}.{component}"] = book(lay.result_cell(entity, component))
    return out


def _app_inplane(
    case: SceneCase, mutate: Callable[[dict[str, Any]], None] | None = None
) -> dict[str, float] | None:
    """The app's rope tensions and reaction components (in-plane and out-of-plane) for a
    (mutated) 22689 document, straight from ``CraneCalculator.compute_analysis_point``
    (``None`` if it finds none), read as the Scene results table reads them."""
    from crane_statics.calculation.engine import CraneCalculator
    from crane_statics.core.demand_extraction import extract_connection_component_response
    from crane_statics.data.models import CraneDocument

    doc_dict = load_document(YAML_PATH).model_dump()
    if mutate is not None:
        mutate(doc_dict)
    doc = CraneDocument.model_validate(doc_dict)
    diameter = doc.support_widths.slewing_bearing_diameter_m
    kwargs = dataclasses.asdict(case)
    hoist_id = kwargs.pop("hoist_id")
    angle, hook = kwargs.pop("boom_angle_deg"), kwargs.pop("hook_load_kN")
    pt = CraneCalculator(doc).compute_analysis_point(angle, hook, hoist_id, **kwargs)
    if pt is None:
        return None
    solve = {r["id"] for r in doc_dict["ropes"] if r.get("tension_mode") == "solve"}
    out: dict[str, float] = {}
    for hoist in doc_dict["hoists"]:
        channel = "boom_hoist" if hoist["rope_id"] in solve else hoist["rope_id"]
        base = pt.baseline_rope_tensions.get(channel, 0.0)
        sens = pt.sensitivity_rope_tensions.get(channel, 0.0)
        out[f"tension.{hoist['id']}"] = base + sens * hook
    for entity, component in layout.DEFAULT_LAYOUT.result_rows:
        comp = {"fx": "Fx", "fy": "Fy", "fz": "Fz", "mx": "Mx", "my": "My"}.get(
            component, component
        )
        value = extract_connection_component_response(
            pt, hook, entity, comp, support_widths_diameter_m=diameter
        )
        assert value is not None, (entity, component)
        if component == "min_bearing_line_load":
            value = max(0.0, value)  # the results row shows uplift only
        out[f"reaction.{entity}.{component}"] = value
    return out


def _solver_statics(
    case: SceneCase, mutate: Callable[[dict[str, Any]], None] | None = None
) -> dict[str, float]:
    """``crane_solver._run_statics`` at ``case``, with the inputs ``_solver_loads`` uses."""
    doc = load_document(YAML_PATH).model_dump()
    if mutate is not None:
        mutate(doc)
    positions = active_hook_positions(
        doc,
        case.hoist_id,
        KinematicsAdapter(doc).compute_all_node_positions(case.boom_angle_deg),
        case.tackle_length_m,
    )
    motion = crane_frame_motion(
        abs(float(doc["environment"]["gravity_m_s2"])),
        VesselInputs(
            accel_x_m_s2=case.accel_x_m_s2,
            accel_y_m_s2=case.accel_y_m_s2,
            accel_z_m_s2=case.accel_z_m_s2,
            roll_deg=case.vessel_roll_deg,
            pitch_deg=case.vessel_pitch_deg,
            slew_deg=case.crane_slewing_angle_deg,
            roll_accel_rad_s2=case.roll_accel_rad_s2,
            pitch_accel_rad_s2=case.pitch_accel_rad_s2,
        ),
    )
    rope_mass = {
        nid: entry.mass_t
        for nid, entry in rope_mass_at_nodes(
            doc, positions, effective_mult=falls_override_effective_mult(doc, {})
        ).items()
    }
    wind_x, _wind_y = split_wind_speed_components(case.wind_speed_m_s, case.wind_direction)
    raw: dict[str, float] = crane_solver._run_statics(
        doc,
        {nid: (float(p["x"]), float(p["z"])) for nid, p in positions.items()},
        {n["id"]: n.get("body", "") for n in doc["nodes"]},
        case.hoist_id,
        case.hook_load_kN,
        wind_x,
        motion.accel_x_m_s2,
        motion.accel_z_m_s2,
        motion.pitch_accel_rad_s2,
        offlead_lateral_force_kN=math.sin(math.radians(case.offlead_deg)) * case.hook_load_kN,
        solver_ctx=crane_solver.build_solver_static_context(doc),
        rope_mass_t_by_node=rope_mass,
        boom_angle_deg=case.boom_angle_deg,
    )
    return raw


def _assert_statics_match_solver(
    book: "_Book", lay: layout.Layout, raw: dict[str, float], tag: Any
) -> None:
    assert set(raw) == set(layout.STATICS_UNKNOWNS), tag
    for key, want in raw.items():
        got = book(lay.statics_value(key))
        assert math.isclose(got, want, rel_tol=1e-9, abs_tol=1e-6), (tag, key, got, want)


def _collapse_rear_leg(model: dict[str, Any]) -> None:
    """Move the rear hinge onto the upper hinge: the rear leg has no length."""
    points = model["structure"]["points"]
    points["rear_hinge"]["x_m"] = points["upper_hinge"]["x_m"]
    points["rear_hinge"]["z_m"] = points["upper_hinge"]["z_m"]


def _collapse_rear_leg_doc(doc: dict[str, Any]) -> None:
    nodes = {n["id"]: n for n in doc["nodes"]}
    for node_id in ("rear_base", "platform_rear"):
        nodes[node_id]["x_m"] = nodes["rear_top"]["x_m"]
        nodes[node_id]["z_m"] = nodes["rear_top"]["z_m"]


def _assert_results_show_message(path: Path, lay: layout.Layout, words: str) -> "_Book":
    _assert_no_errors(path, layout.RESULTS)
    book = _Book(path)
    assert book(lay.statics_cell("statics_status", "ok")) == 0
    message = book(lay.results_cell("results_status", "message"))
    assert isinstance(message, str) and words in message.lower(), message
    for entity, component in lay.result_rows:
        assert isinstance(book(lay.result_cell(entity, component)), str), (entity, component)
    for i in range(1, len(_model()["hoists"]) + 1):
        assert isinstance(book(lay.results_addr("tensions", "tension_kN", i)), str), i
    return book


#: Golden out-of-plane components the workbook must reproduce. heel_left/heel_right
#: boom_axial/boom_normal need the name-based heel axis node and are omitted, as for the
#: heel point.
OOP_REQUIRED = {
    ("heel_left", "fx"),
    ("heel_left", "fz"),
    ("heel_left", "resultant"),
    ("heel_right", "fx"),
    ("heel_right", "fz"),
    ("heel_right", "resultant"),
    ("platform_base", "fy"),
    ("platform_base", "mx"),
    ("platform_base", "fradial"),
    ("platform_base", "mradial"),
    ("platform_base", "max_bearing_line_load"),
    ("platform_base", "min_bearing_line_load"),
}


#: Support group of the engine's out-of-plane pass -> its (left, right) split entities.
OOP_SIDES = {
    "heel": ("heel_left", "heel_right"),
    "front_leg": ("front_leg_left", "front_leg_right"),
    "rear_leg": ("rear_leg_left", "rear_leg_right"),
}


def _assert_book_clean(book: "_Book", *sheets: str) -> None:
    for sheet in sheets:
        values = book.values[sheet]
        errors = {c: v for c, v in values.items() if isinstance(v, str) and v.startswith("#")}
        assert not errors, (sheet, list(errors.items())[:10])


def _case_motion(case: SceneCase) -> Any:
    return crane_frame_motion(
        abs(float(_model()["environment"]["gravity_m_s2"])),
        VesselInputs(
            accel_x_m_s2=case.accel_x_m_s2,
            accel_y_m_s2=case.accel_y_m_s2,
            accel_z_m_s2=case.accel_z_m_s2,
            roll_deg=case.vessel_roll_deg,
            pitch_deg=case.vessel_pitch_deg,
            slew_deg=case.crane_slewing_angle_deg,
            roll_accel_rad_s2=case.roll_accel_rad_s2,
            pitch_accel_rad_s2=case.pitch_accel_rad_s2,
        ),
    )


def _engine_oop(case: SceneCase) -> dict[str, Any]:
    """The engine's out-of-plane pass at ``case``, called directly.

    ``CraneCalculator.compute_analysis_point`` runs it (and leaves its point geometry, with
    the lumped hook and rope masses, in ``_oop_geometry``); ``_out_of_plane_variants`` is then
    evaluated once with the actual hook load and side force ``sin(sidelead) x hook load``.
    The group moments are recovered from the split: ``M = (left - right) x width / 2``.
    """
    from crane_statics.calculation.engine import CraneCalculator

    doc = load_document(YAML_PATH)
    calc = CraneCalculator(doc)
    kwargs = dataclasses.asdict(case)
    hoist_id = kwargs.pop("hoist_id")
    angle, hook = kwargs.pop("boom_angle_deg"), kwargs.pop("hook_load_kN")
    pt = calc.compute_analysis_point(angle, hook, hoist_id, **kwargs)
    assert pt is not None and calc._oop_geometry is not None
    (_angle, lumped, _rho), geometry = calc._oop_geometry
    motion = _case_motion(case)
    _wind_x, wind_y = split_wind_speed_components(case.wind_speed_m_s, case.wind_direction)
    side_kN = math.sin(math.radians(case.sidelead_deg)) * hook
    ((split, platform_mx, platform_fy),) = calc._out_of_plane_variants(
        pt,
        hoist_id,
        angle,
        [(hook, side_kN)],
        wind_y,
        motion.accel_y_m_s2,
        motion.roll_accel_rad_s2,
        accel_x_m_s2=motion.accel_x_m_s2,
        accel_z_m_s2=motion.accel_z_m_s2,
        pitch_accel_rad_s2=motion.pitch_accel_rad_s2,
        lumped_point_masses=lumped,
    )
    width = float(doc.support_widths.boom_width_m)
    moments: dict[str, dict[str, float]] = {
        group: {
            "mx": (split[left]["Fz"] - split[right]["Fz"]) * width / 2.0,
            "mz": (split[left]["Fx"] - split[right]["Fx"]) * width / 2.0,
        }
        for group, (left, right) in OOP_SIDES.items()
    }
    moments["platform_base"] = {"mx": platform_mx}
    assert doc.environment.vessel_motion is not None
    ref_z = float(doc.environment.vessel_motion.reference_z_m)
    # _out_of_plane_variants' fy_m and fy_w, on the engine's own point geometry.
    fy = -geometry.mass_t * (
        motion.accel_y_m_s2 - motion.roll_accel_rad_s2 * (geometry.mass_z - ref_z)
    )
    n_doc = sum(len(body.mass_points) for body in doc.bodies)
    sign = 0.0 if abs(wind_y) < 1e-12 else math.copysign(1.0, wind_y)
    return {
        "split": split,
        "moments": moments,
        "platform_fy": platform_fy,
        "mass_fy": fy[:n_doc],
        "mass_y": geometry.mass_y[:n_doc],
        "lumped_fy": float(fy[n_doc:].sum()),
        "wind_fy": sign * geometry.wind_coef * abs(wind_y) ** 2,
        "side_kN": side_kN,
    }


def _assert_oop_matches_engine(
    book: "_Book", lay: layout.Layout, model: dict[str, Any], case: SceneCase, tag: Any
) -> None:
    want = _engine_oop(case)
    oop = lay.oop

    def close(got: Any, value: float, what: Any) -> None:
        assert math.isclose(got, value, rel_tol=1e-9, abs_tol=1e-6), (tag, what, got, value)

    assert len(want["mass_fy"]) == sum(len(b["mass_points"]) for b in model["bodies"])
    for i, value in enumerate(want["mass_fy"], start=1):
        close(book(oop.addr("oop_masses", "fy_kN", i)), value, ("mass fy", i))
        close(book(oop.addr("oop_masses", "y_m", i)), want["mass_y"][i - 1], ("mass y", i))
    for i, value in enumerate(want["wind_fy"], start=1):
        close(book(oop.addr("oop_winds", "fy_kN", i)), value, ("wind fy", i))
    lumped = sum(
        book(oop.addr("oop_lumped", "fy_kN", i))
        for i in range(1, oop.tables["oop_lumped"].capacity + 1)
    )
    close(lumped, want["lumped_fy"], "lumped fy")
    close(book(oop.cell("sidelead", "fy_kN")), want["side_kN"], "side force")
    for group, comps in want["moments"].items():
        for comp, value in comps.items():
            close(book(oop.moment(group, comp)), value, (group, comp))
    for entity, comps in want["split"].items():
        for comp, value in comps.items():
            close(book(oop.split_value(entity, comp.lower())), value, (entity, comp))
    close(book(oop.cell("platform", "fy_kN")), want["platform_fy"], "platform fy")
    close(book(oop.cell("platform", "mx_kNm")), want["moments"]["platform_base"]["mx"], "mx")


#: Golden keys the workbook deliberately does not produce: the heel's boom_axial and
#: boom_normal projections need the app's name-based heel axis node (reference.py rule 6).
EXPECTED_OMISSIONS = {
    f"reaction.{entity}.{component}"
    for entity in ("heel_point", "heel_left", "heel_right")
    for component in ("boom_axial", "boom_normal")
}


#: Acceptance tolerance: 0.01 kN for forces, 0.01 kNm for moments and 0.01 kN/m for the
#: platform bearing line loads (the same absolute figure in each row's own unit).
ACCEPT_TOL = 0.01


DELIVERED = Path(__file__).resolve().parents[1] / "docs" / "crane-statics-validation-22689.xlsx"


def case_book_factory(tmp_path_factory: pytest.TempPathFactory) -> Callable[[str], _Book]:
    """``case_book(name)``: the recalculated default-layout workbook of an unmodified 22689
    case, built and recalculated once per session (per xdist worker). Read only: tests that
    change a workbook build their own."""
    books: dict[str, _Book] = {}

    def get(name: str) -> _Book:
        if name not in books:
            path = tmp_path_factory.mktemp(name) / f"case_{name}.xlsx"
            build_workbook(_model(), CASES[name], path)
            books[name] = _Book(path)
        return books[name]

    return get
