"""Tests for the Excel validation workbook tooling."""

import dataclasses
import json
import math
from pathlib import Path
from typing import Any

import numpy as np
import pytest
from openpyxl import Workbook
from openpyxl.worksheet.formula import ArrayFormula

from scripts.excel_workbook.recalc import recalculate
from scripts.excel_workbook.reference import (
    CASES,
    GOLDEN_PATH,
    YAML_PATH,
    app_results,
    expanded_model,
    scene_reaction_arrows,
)

CASE_NAMES = {"base", "inplane_wind_accel", "outofplane_wind_roll_sidelead", "idle_hoist"}


def _num(value: float | str | None) -> float:
    assert isinstance(value, float), value
    return value


def test_recalculate_supports_minverse_mmult(tmp_path: Path) -> None:
    matrix = [[4.0, -2.0, 1.0], [3.0, 6.0, -4.0], [2.0, 1.0, 8.0]]
    vector = [12.0, -25.0, 32.0]
    wb = Workbook()
    ws = wb.active
    ws.title = "Solve"
    for r, row in enumerate(matrix, start=1):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=value)
        ws.cell(row=r, column=5, value=vector[r - 1])
    ws["G1"] = ArrayFormula("G1:G3", "=MMULT(MINVERSE(A1:C3),E1:E3)")
    path = tmp_path / "solve.xlsx"
    wb.save(path)

    values = recalculate(path)["Solve"]

    expected = np.linalg.solve(np.array(matrix), np.array(vector))
    got = [values["G1"], values["G2"], values["G3"]]
    assert np.allclose(got, expected, rtol=0, atol=1e-9)


def test_recalculate_reads_sumproduct_if_index_match(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Fn"
    for i, (key, val) in enumerate([(1, 10.0), (2, 20.0), (3, 30.0)], start=1):
        ws.cell(row=i, column=1, value=key)
        ws.cell(row=i, column=2, value=val)
    ws["D1"] = "=SUMPRODUCT(A1:A3,B1:B3)"
    ws["D2"] = '=IF(B1>5,"big","small")'
    ws["D3"] = "=INDEX(B1:B3,MATCH(2,A1:A3,0))"
    ws["D4"] = "=ABS(-3.5)"
    ws["D5"] = "=SQRT(16)"
    ws["D6"] = "=SIN(0.5)"
    ws["D7"] = "=COS(0.5)"
    ws["D8"] = "=ATAN2(1,2)"
    path = tmp_path / "fn.xlsx"
    wb.save(path)

    values = recalculate(path)["Fn"]

    assert values["D1"] == 140.0
    assert values["D2"] == "big"
    assert values["D3"] == 20.0
    assert values["D4"] == 3.5
    assert values["D5"] == 4.0
    assert math.isclose(_num(values["D6"]), math.sin(0.5), abs_tol=1e-12)
    assert math.isclose(_num(values["D7"]), math.cos(0.5), abs_tol=1e-12)
    # Excel ATAN2(x, y) takes x first.
    assert math.isclose(_num(values["D8"]), math.atan2(2, 1), abs_tol=1e-12)


def test_recalculate_extra_functions(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Extra"
    for i, (k, v) in enumerate([("a", 1.0), ("b", 2.0), ("a", 4.0)], start=1):
        ws.cell(row=i, column=1, value=k)
        ws.cell(row=i, column=2, value=v)
    formulas = {
        "D1": ("=RADIANS(180)", math.pi),
        "D2": ("=DEGREES(PI())", 180.0),
        "D3": ("=ACOS(0.5)", math.acos(0.5)),
        "D4": ("=ATAN(1)", math.atan(1)),
        "D5": ("=MAX(B1:B3)", 4.0),
        "D6": ("=MIN(B1:B3)", 1.0),
        "D7": ("=SUM(B1:B3)", 7.0),
        "D8": ('=SUMIF(A1:A3,"a",B1:B3)', 5.0),
        "D9": ('=SUMIFS(B1:B3,A1:A3,"a",B1:B3,">1")', 4.0),
        "D10": ("=SQRT(3^2+4^2)", 5.0),
    }
    for cell, (formula, _) in formulas.items():
        ws[cell] = formula
    path = tmp_path / "extra.xlsx"
    wb.save(path)

    values = recalculate(path)["Extra"]

    for cell, (_, expected) in formulas.items():
        assert math.isclose(_num(values[cell]), expected, abs_tol=1e-12), cell


def _golden() -> dict[str, Any]:
    golden: dict[str, Any] = json.loads(GOLDEN_PATH.read_text(encoding="utf-8"))
    return golden


def test_golden_cases_match_live_app() -> None:
    golden = _golden()
    assert set(golden["cases"]) == CASE_NAMES
    assert set(CASES) == CASE_NAMES
    for name, entry in golden["cases"].items():
        assert entry["inputs"] == dataclasses.asdict(CASES[name]), name
        live = app_results(YAML_PATH, CASES[name])
        stored = entry["app_results"]
        assert set(stored) == set(live), name
        for key, value in stored.items():
            assert math.isclose(value, live[key], rel_tol=0.0, abs_tol=1e-9), (name, key)


@pytest.mark.parametrize("name", sorted(CASE_NAMES))
def test_reaction_arrows_match_result_rows(name: str) -> None:
    results = app_results(YAML_PATH, CASES[name])
    arrows = scene_reaction_arrows(YAML_PATH, CASES[name])
    assert set(arrows) == {
        "heel_point",
        "rear_hinge",
        "front_hinge",
        "upper_hinge",
        "platform_base",
    }
    for entity, comps in arrows.items():
        for comp, value in comps.items():
            assert math.isclose(
                value, results[f"reaction.{entity}.{comp}"], rel_tol=0.0, abs_tol=1e-9
            ), (name, entity, comp)


def test_cases_exercise_their_feature() -> None:
    base, wind, oop, idle = (
        CASES[k]
        for k in ("base", "inplane_wind_accel", "outofplane_wind_roll_sidelead", "idle_hoist")
    )
    assert base.hook_load_kN > 0 and base.wind_speed_m_s == 0 and base.accel_x_m_s2 == 0
    assert wind.wind_direction in {"+X", "-X"} and wind.wind_speed_m_s > 0
    assert wind.accel_x_m_s2 != 0 and wind.accel_z_m_s2 != 0 and wind.pitch_accel_rad_s2 != 0
    assert oop.wind_direction in {"+Y", "-Y"} and oop.wind_speed_m_s > 0
    assert oop.vessel_roll_deg != 0 and oop.sidelead_deg != 0 and oop.roll_accel_rad_s2 != 0
    assert idle.hoist_id != base.hoist_id
    assert idle.vessel_roll_deg == 0 and idle.vessel_pitch_deg == 0
    oop_results = app_results(YAML_PATH, oop)
    assert abs(oop_results["reaction.platform_base.fy"]) > 1.0
    assert abs(oop_results["reaction.platform_base.mx"]) > 1.0
    # The base case's hook-loaded hoist is idle in the idle-hoist case: its rope
    # carries only its hook block, m_hook * (g + az) / falls (no vessel tilt here).
    idle_results = app_results(YAML_PATH, idle)
    model = expanded_model(YAML_PATH)
    mh = next(h for h in model["hoists"] if h["id"] == base.hoist_id)
    g = abs(model["environment"]["gravity_m_s2"])
    expected = mh["hook_mass_t"] * (g + idle.accel_z_m_s2) / mh["falls"]
    assert math.isclose(idle_results[f"tension.{base.hoist_id}"], expected, abs_tol=1e-9)


def test_expanded_model_is_json_serialisable() -> None:
    model = expanded_model(YAML_PATH)
    json.dumps(model)
    for key in ("environment", "structure", "bodies", "nodes", "ropes", "hoists"):
        assert key in model
